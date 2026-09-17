# -*- coding: utf-8 -*-
"""阶段 3：入选跟踪结果导出（openpyxl 直写 xlsx）。

口径与产物定义见 docs/plans/auto-screen-track/contract.md §3/§9/§10。
只读导出的数据源与只读路由（api_routes/tracking.py）完全一致：
只认「发布索引指向的快照」+「当前 revision 的 track 产物」（契约 §0/§5），
不重算跟踪、不写快照/产物。

三个数据 sheet：
- 指标汇总：一行一个 (rule_id, rule_fingerprint) 分段；
- 周汇总：一行一个 (规则, 发布周)；
- 周明细：一行一只票 (规则, 周, 代码)。

关键取舍（为什么这样做）：
- **空值写空单元格，不写 0**：契约 §3 规定空仓周胜率/覆盖率为 null。
  0 是有意义的数值（"零收益/零覆盖"），与"无有效样本"语义相反；写 0
  会让使用者误读，所以 None 一律保留 None。
- **收益率/胜率统一换算成百分数（×100）** 并在表头标注「%」，三个 sheet
  口径一致，避免明细是 5.00、汇总却是 0.05 的割裂读感。
- **回填提示只取一份文案**：从 api_routes/tracking.py 的 BACKFILL_NOTICE
  取（函数内延迟 import，避免 service→api 的模块级耦合），不另写措辞，
  防止两处文案漂移。
- 定位/读取发布快照与当前产物的私有帮助函数直接复用 tracking.py 的实现，
  不复制第二份定位逻辑（契约 §0 指针是唯一事实源）。
"""

from __future__ import annotations

import json
import time
import uuid
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from . import screen_contract as sc

# 固定 sheet 名（不含用户输入，天然不会冲突/超长，无需 sanitize 消解）
SHEET_SUMMARY = "指标汇总"
SHEET_WEEK = "周汇总"
SHEET_DETAIL = "周明细"
SHEET_META = "meta"
SHEETS_ORDER = (SHEET_SUMMARY, SHEET_WEEK, SHEET_DETAIL, SHEET_META)

# 导出落盘目录（与卦象导出一致：storage/astock/bagua_exports）
EXPORT_SUBDIRNAME = "bagua_exports"

# 成交性中文展示（契约 §1：绝不默认"正常"，未知如实标注）
FILL_LABELS = {
    sc.FILL_OK: "可成交",
    sc.FILL_LIMIT_UP_UNBUYABLE: "一字涨停买不进",
    sc.FILL_NO_BAR: "无K线",
    sc.FILL_UNKNOWN: "判定未知",
}

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

# YYYYMMDD → 中文星期：导出「见顶星期」列与 UI 列表同口径（数据同源，
# 只是把产物里的日期转成人读的星期）。
_WEEKDAYS_CN = ("周一", "周二", "周三", "周四", "周五", "周六", "周日")


def _weekday_cn(date_int: Any) -> str:
    """见顶日 → 中文星期；缺失/非法日期留空串（不猜、不填占位值）。"""
    try:
        from datetime import date as _date

        d = int(date_int)
        return _WEEKDAYS_CN[_date(d // 10000, (d // 100) % 100, d % 100).weekday()]
    except Exception:  # noqa: BLE001
        return ""


def _current_rule_catalog(cfg) -> Optional[Dict[str, Dict[str, Any]]]:
    """当前规则目录（规则中心可见 ∪ 周五链预置复核规则）：id -> 判定信息。

    2026-09-15 用户要求「跟踪与规则中心一致」：导出必须与页面同口径——
    规则中心删掉的规则不再导出、同一公式的多个 rule_id 归并成一行，否则
    导出的行数与页面显示对不上账（契约要求 UI 与导出一致）。

    目录本体在 service/screening.current_rule_catalog（与跟踪页/只读 API
    共用同一实现，保证两处选出同一个 canonical）。本函数只做异常降级：
    返回 None 表示目录不可得（规则库异常）→ 调用方不做过滤与归并，
    宁可多导出历史数据，也不能让导出直接失败。
    """
    try:
        from .screening import current_rule_catalog

        return current_rule_catalog(cfg)
    except Exception:  # noqa: BLE001 — 目录不可用 → 不过滤（fail-open）
        return None


def _pick_canonical(
    ids: List[str], catalog: Optional[Dict[str, Dict[str, Any]]]
) -> str:
    """同指纹多 id 里选代表身份。

    打分与页面共用 service/screening.pick_canonical_rule_id 的同一实现
    （目录内存在 > 可执行 > 未隐藏 > 来源 user>builtin>system > id 字典序）——
    两处算法若不同，导出与页面可能对同一组选出不同代表，行数就对不上账了。
    """
    from .screening import pick_canonical_rule_id

    return pick_canonical_rule_id(ids, catalog)


# ---------------------------------------------------------------------------
# Excel 单元格清洗：与 service/bagua_query.py 的 _excel_safe_cell 同源（复制）
#
# 该实现引用的 bagua_query 模块会拉起 numpy/data_store 等重依赖（import 约 1s），
# 为一个纯格式小函数把整条重依赖链拖进跟踪导出不值得，故按其原文复制并注明来源，
# 行为保持逐字一致（控制字符替换、公式注入前缀单引号、32767 长度截断）。
# ---------------------------------------------------------------------------
_EXCEL_FORMULA_LEAD = ("=", "+", "-", "@")  # 来源：bagua_query._EXCEL_FORMULA_LEAD
_EXCEL_ILLEGAL_CHARS = frozenset(  # 来源：bagua_query._EXCEL_ILLEGAL_CHARS
    [chr(c) for c in range(0x00, 0x09)]
    + [chr(c) for c in range(0x0B, 0x0D)]
    + [chr(c) for c in range(0x0E, 0x20)]
    + [chr(c) for c in range(0x7F, 0xA0)]
    + ["\ufffe", "\uffff"]
) | frozenset(chr(c) for c in range(0xD800, 0xE000))
_EXCEL_MAX_CELL_LEN = 32767


def _excel_safe_cell(value: Any) -> Any:
    """清洗 + 防 Excel 公式注入（复制自 service/bagua_query.py，来源见上）。

    字符串先移除/替换 Excel 与 XML 非法字符；以 ``=+-@`` 开头的加前缀单引号
    使其落成普通文本而非公式；None/数值原样返回（None → 空单元格）。
    """
    if isinstance(value, str):
        if any(ch in _EXCEL_ILLEGAL_CHARS for ch in value):
            value = "".join(
                "_" if ch in _EXCEL_ILLEGAL_CHARS else ch for ch in value
            )
        if value[:1] in _EXCEL_FORMULA_LEAD:
            value = "'" + value[:_EXCEL_MAX_CELL_LEN - 1]
        elif len(value) > _EXCEL_MAX_CELL_LEN:
            value = value[:_EXCEL_MAX_CELL_LEN]
    return value


def _pct(value: Any) -> Optional[float]:
    """比率 → 百分数数值（0.05 → 5.0）；None 保持 None（空单元格）。"""
    if value is None:
        return None
    try:
        return round(float(value) * 100.0, 4)
    except (TypeError, ValueError):
        return None


def _backfill_notice() -> str:
    """契约 §9 回填提示全文：单一文案来源（api_routes/tracking.py）。

    函数内延迟 import：本模块被 tracking.py 的新端点调用，模块级 import
    会形成 service→api_routes 的加载期耦合；延迟到调用时双方均已就绪。
    """
    from ..api_routes.tracking import BACKFILL_NOTICE

    return BACKFILL_NOTICE


def _mean(values: List[Any]) -> Optional[float]:
    vals = [v for v in values if v is not None]
    if not vals:
        return None
    return sum(vals) / len(vals)


def _uniq_json(items: List[Any]) -> str:
    """唯一化后序列化（1 个写单值 JSON，多个写 JSON 数组），便于 meta 追溯。"""
    seen: List[str] = []
    out: List[Any] = []
    for it in items:
        s = json.dumps(it, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        if s not in seen:
            seen.append(s)
            out.append(it)
    if not out:
        return ""
    if len(out) == 1:
        return seen[0]
    return json.dumps(out, ensure_ascii=False, sort_keys=True)


def _agg_get(agg: Optional[dict], *keys: str) -> Any:
    """按顺序取第一个非 None 的字段值（兼容契约名与实现名的历史差异）。"""
    if not agg:
        return None
    for k in keys:
        v = agg.get(k)
        if v is not None:
            return v
    return None


def _mean_max_gain(track: Optional[dict], rule_id: str, agg: Optional[dict]) -> Optional[float]:
    """平均周内最大涨幅：优先聚合字段，缺失时从 rows 现算。

    （阶段 2 实现产出的是 mean_giveback / 无 mean_max_gain_sig，与契约 §10
    列名不完全一致，这里两种名字都容忍，避免导出依赖某一种。）
    """
    v = _agg_get(agg, "mean_max_gain_sig")
    if v is not None:
        return v
    vals = [
        r.get("max_gain_sig")
        for r in (track or {}).get("rows") or []
        if str(r.get("rule_id")) == rule_id and r.get("max_gain_sig") is not None
    ]
    return _mean(vals)


def _mean_giveback(track: Optional[dict], rule_id: str, agg: Optional[dict]) -> Optional[float]:
    """平均回吐：优先聚合字段（兼容两种名字），缺失时按 max_gain-ret_close 现算。"""
    v = _agg_get(agg, "mean_giveback_sig", "mean_giveback")
    if v is not None:
        return v
    vals = [
        r["max_gain_sig"] - r["ret_close_sig"]
        for r in (track or {}).get("rows") or []
        if str(r.get("rule_id")) == rule_id
        and r.get("max_gain_sig") is not None
        and r.get("ret_close_sig") is not None
    ]
    return _mean(vals)


def _write_sheet(ws, headers: List[str], rows: List[List[Any]], notice: Optional[str]) -> int:
    """写一个数据 sheet：可选回填提示置顶，再写表头与数据行。

    返回数据行数。回填提示只在该 sheet 含回填周时传入（见调用方）。
    """
    from openpyxl.styles import Font

    row_idx = 1
    if notice:
        ws.cell(row_idx, 1, _excel_safe_cell(notice)).font = Font(bold=True)
        row_idx += 2  # 提示与表头之间留一空行
    header_row = row_idx
    for ci, h in enumerate(headers, 1):
        ws.cell(header_row, ci, _excel_safe_cell(h)).font = Font(bold=True)
    row_idx += 1
    for row in rows:
        for ci, v in enumerate(row, 1):
            ws.cell(row_idx, ci, _excel_safe_cell(v))
        row_idx += 1
    # 冻结表头下方：滚动查看明细时列名不丢
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    return len(rows)


def export_tracking_xlsx(
    cfg,
    *,
    weeks: int = 12,
    rule_id: Optional[str] = None,
    entry_asof: Optional[int] = None,
) -> Dict[str, Any]:
    """把已结算的跟踪产物导出为 xlsx，返回结果摘要（不抛业务异常）。

    weeks      : 取最近 N 个发布信号自然周（周历序，非"有收益的周"）；
    rule_id    : 只导出该规则（缺省=全部）；
    entry_asof : 限定单周导出（8 位信号日）；给了它则忽略 weeks 窗口。

    无任何发布周 → {"ok": False, "reason": "no_published_week"}（不产空文件）。
    """
    # 延迟 import：复用 tracking.py 的指针定位/读取实现（唯一事实源），
    # 同时避免模块级 service→api_routes 循环耦合。
    from ..api_routes import tracking as tr
    from .screen_snapshots import load_published_snapshot_for_week

    rule_filter = (str(rule_id).strip() if rule_id is not None else "") or None

    week_map = tr._published_weeks(cfg)
    if entry_asof is not None:
        wid = int(entry_asof)
        week_ids = [wid] if str(wid) in week_map else []
    else:
        all_ids = sorted(int(k) for k in week_map.keys())
        week_ids = all_ids[-int(weeks):] if int(weeks) > 0 else []
    if not week_ids:
        return {"ok": False, "reason": "no_published_week"}

    # 逐周读发布快照 + 当前 revision 产物；快照读不到（索引坏了）跳过
    per_week: List[Dict[str, Any]] = []
    for wid in week_ids:
        snap = load_published_snapshot_for_week(cfg, wid)
        if snap is None:
            continue
        track = tr._track_for_week(cfg, snap, wid)
        run_kind = str(snap.get("run_kind") or (track or {}).get("run_kind") or "")
        per_week.append(
            {
                "wid": wid,
                "snap": snap,
                "track": track,
                "run_kind": run_kind,
                "backfill": run_kind == "backfill",
                # 规则范围（2026-09-16）：subset = 「指定规则补算」周。导出必须
                # 如实标注——该周只有这几条规则跑过筛选，其余规则当周无名单，
                # 不标会被读成"其他规则当周空仓"（meta 里逐周列明）。
                "scope": sc.snapshot_rules_scope(snap),
                "scope_ids": sc.scoped_rule_ids(snap),
            }
        )
    if not per_week:
        return {"ok": False, "reason": "no_published_week"}

    def _agg_for(track: Optional[dict], rid: str) -> Optional[dict]:
        if not track:
            return None
        for a in track.get("rule_aggregates") or []:
            if str(a.get("rule_id")) == rid:
                return tr._with_exec_excess_stats(track, rid, a)
        return None

    # ---- 目录过滤 + 同指纹归并（2026-09-15：与跟踪页/只读 API 同口径）----
    # 规则中心已删除的规则不出现在导出里；同一公式的多个 rule_id（tn6_ 与其
    # 配对源 txt_、改名前后的 user_ 规则）归并成一行——否则导出的行数与页面
    # 对不上账（契约：UI 与导出必须同口径）。
    catalog = _current_rule_catalog(cfg)
    fp_group: Dict[str, set] = {}
    for pw in per_week:
        fps_w = pw["snap"].get("rule_fingerprints") or {}
        for r in pw["snap"].get("rules") or []:
            rid_w = str(r.get("rule_id"))
            if catalog is not None and rid_w not in catalog:
                continue
            fp_w = str(fps_w.get(rid_w) or "")
            if fp_w:
                fp_group.setdefault(fp_w, set()).add(rid_w)
    canonical_by_fp: Dict[str, str] = {}
    canonical_by_id: Dict[str, str] = {}
    for fp_w, ids_w in fp_group.items():
        canon_w = _pick_canonical(sorted(ids_w), catalog)
        canonical_by_fp[fp_w] = canon_w
        for i_w in ids_w:
            canonical_by_id[i_w] = canon_w

    def _canon(rid: str, fp: str) -> str:
        """该 (rule_id, 指纹) 归属的 canonical id（无指纹段按自身）。"""
        if fp and fp in canonical_by_fp:
            return canonical_by_fp[fp]
        return canonical_by_id.get(rid) or rid

    def _hit_filter(rid: str, fp: str) -> bool:
        """规则筛选命中判定：命中同指纹组任一 id 即算命中（与 L2 同口径）。"""
        if rule_filter is None:
            return True
        return rid == rule_filter or _canon(rid, fp) == _canon(rule_filter, "")

    segments: Dict[Tuple[str, str], Dict[str, Any]] = {}
    week_rows: List[Dict[str, Any]] = []
    detail_rows: List[Dict[str, Any]] = []
    warnings: List[str] = []
    unsettled_weeks = 0
    summary_backfill = False
    week_backfill = False
    detail_backfill = False

    for pw in per_week:
        wid = pw["wid"]
        snap = pw["snap"]
        track = pw["track"]
        is_backfill = pw["backfill"]
        if track is None:
            unsettled_weeks += 1
        fps = snap.get("rule_fingerprints") or {}
        completion = (track or {}).get("completion") or "no_product"
        settled = bool(track and completion == sc.TRACK_COMPLETE)
        for r in snap.get("rules") or []:
            rid = str(r.get("rule_id"))
            if catalog is not None and rid not in catalog:
                continue  # 规则中心已删除的规则：删除同步（不导出）
            fp = str(fps.get(rid) or "")
            if not _hit_filter(rid, fp):
                continue
            canon = _canon(rid, fp)
            agg = _agg_for(track, rid)
            selected = (
                int(agg.get("selected_count"))
                if agg is not None and agg.get("selected_count") is not None
                else int(r.get("count") or 0)
            )
            seg = segments.setdefault(
                (canon, fp), {"rule_id": canon, "fingerprint": fp, "weeks": []}
            )
            if any(w["wid"] == wid for w in seg["weeks"]):
                # 同周同指纹组已有数据（兄弟 id 同公式、数据等价）：只取一份，
                # 否则选票数/收益会被双计
                continue
            seg["weeks"].append(
                {"wid": wid, "settled": settled, "selected": selected, "agg": agg, "backfill": is_backfill}
            )
            week_rows.append(
                {
                    "rid": canon,
                    "wid": wid,
                    "asof": snap.get("asof"),
                    "run_kind": pw["run_kind"],
                    "selected": selected,
                    "agg": agg,
                    "track": track,
                    "completion": completion,
                    "backfill": is_backfill,
                }
            )
            if is_backfill:
                summary_backfill = True
                week_backfill = True
        seen_detail: set = set()
        for row in (track or {}).get("rows") or []:
            rid = str(row.get("rule_id"))
            if catalog is not None and rid not in catalog:
                continue  # 已删规则的明细不导出
            fp_of_row = str(fps.get(rid) or "")
            if not _hit_filter(rid, fp_of_row):
                continue
            canon = _canon(rid, fp_of_row)
            # 同一周同一票在归并组内只留一行（规则 ID 列统一显示 canonical，
            # 与页面一致；原始 id 写在产物里不改）
            key = (wid, canon, str(row.get("code")))
            if key in seen_detail:
                continue
            seen_detail.add(key)
            detail_rows.append(
                {
                    "wid": wid,
                    "row": dict(row, rule_id=canon),
                    "backfill": is_backfill,
                }
            )
            if is_backfill:
                detail_backfill = True

    # ---- 指标汇总（rule × fingerprint 分段，跨周默认每周等权）----
    summary_out: List[List[Any]] = []
    for (rid, fp), seg in sorted(segments.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        seg_weeks = seg["weeks"]
        settled = [w for w in seg_weeks if w["settled"] and w["agg"] is not None]
        weekly_sig = [w["agg"].get("mean_ret_close_sig") for w in settled]
        weekly_exec = [w["agg"].get("mean_ret_close_exec") for w in settled]
        weekly_excess_sig = [w["agg"].get("mean_excess_sig") for w in settled]
        weekly_excess_exec = [w["agg"].get("mean_excess_exec") for w in settled]
        total_selected = sum(w["selected"] or 0 for w in seg_weeks)
        valid_sig_weeks = len([v for v in weekly_sig if v is not None])
        # 周等权胜率：各周 win_rate_sig 的均值（与 L0 API 的
        # weekly_equal_win_rate_sig 同一口径；绝不另用"正收益周占比"——
        # 同名列两种算法会让 UI 与导出 xlsx 数字对不上账）
        weekly_wr_sig = [w["agg"].get("win_rate_sig") for w in settled]
        summary_out.append(
            [
                rid,
                fp,
                len(seg_weeks),
                len(settled),
                total_selected,
                _pct(_mean(weekly_wr_sig)),
                _pct(_mean(weekly_sig)),
                _pct(_mean(weekly_exec)),
                _pct(_mean(weekly_excess_sig)),
                _pct(_mean(weekly_excess_exec)),
                valid_sig_weeks,
                "样本不足" if total_selected < tr.MIN_SAMPLE_WARN else "样本充足",
            ]
        )

    # ---- 周汇总（一行一个 规则×周）----
    week_out: List[List[Any]] = []
    for wr in week_rows:
        agg = wr["agg"]
        week_out.append(
            [
                wr["rid"],
                wr["wid"],
                wr["asof"],
                wr["run_kind"],
                wr["selected"],
                _agg_get(agg, "valid_sig_count"),
                _agg_get(agg, "valid_exec_count"),
                _agg_get(agg, "pending_count"),
                _agg_get(agg, "missing_count"),
                _agg_get(agg, "unbuyable_count"),
                _agg_get(agg, "unknown_count"),
                _pct(_agg_get(agg, "win_rate_sig")),
                _pct(_agg_get(agg, "mean_ret_close_sig")),
                _pct(_agg_get(agg, "mean_ret_close_exec")),
                _pct(_agg_get(agg, "mean_excess_sig")),
                _pct(_agg_get(agg, "mean_excess_exec")),
                _pct(_mean_max_gain(wr["track"], wr["rid"], agg)),
                _pct(_mean_giveback(wr["track"], wr["rid"], agg)),
                wr["completion"],
                "是" if wr["backfill"] else "否",
            ]
        )

    # ---- 周明细（一行一只票；收益率一律百分数）----
    detail_out: List[List[Any]] = []
    for dr in detail_rows:
        r = dr["row"]
        detail_out.append(
            [
                r.get("rule_id"),
                dr["wid"],
                r.get("code"),
                r.get("name") or "",
                r.get("entry_close_signal"),
                r.get("entry_open_week"),
                r.get("close_week_end"),
                FILL_LABELS.get(str(r.get("fill_status")), str(r.get("fill_status") or "")),
                _pct(r.get("max_gain_sig")),
                r.get("max_gain_sig_date"),
                _weekday_cn(r.get("max_gain_sig_date")),
                _pct(r.get("min_low_ret_sig")),
                _pct(r.get("drawdown_close_sig")),
                _pct(r.get("ret_close_sig")),
                _pct(r.get("ret_close_exec")),
                _pct(r.get("theoretical_open_ret")),
                _pct(r.get("bench_ret_sig")),
                _pct(r.get("bench_ret_exec")),
                _pct(r.get("excess_sig")),
                _pct(r.get("excess_exec")),
                r.get("status"),
            ]
        )

    # ---- meta：版本与来源可追溯 + 分口径覆盖率 ----
    tracks = [pw["track"] for pw in per_week if pw["track"]]
    snapshot_ids: List[str] = []
    rev_ids: List[str] = []
    algo_versions: List[str] = []
    schema_versions: List[str] = []
    data_versions: List[Any] = []
    benchmark_versions: List[Any] = []
    for pw in per_week:
        sid = str(pw["snap"].get("snapshot_id") or "")
        if sid and sid not in snapshot_ids:
            snapshot_ids.append(sid)
    for t in tracks:
        rev = str(t.get("tracking_revision_id") or "")
        if rev and rev not in rev_ids:
            rev_ids.append(rev)
        av = str(t.get("algo_version") or "")
        if av and av not in algo_versions:
            algo_versions.append(av)
        sv = str(t.get("schema_version") or "")
        if sv and sv not in schema_versions:
            schema_versions.append(sv)
        data_versions.append(t.get("data_version") or {})
        benchmark_versions.append(t.get("benchmark_data_version") or {})

    tot_sel = tot_sig = tot_exec = tot_exc = 0
    have_agg = False
    for wr in week_rows:
        agg = wr["agg"]
        if agg is None:
            continue
        have_agg = True
        tot_sel += int(agg.get("selected_count") or 0)
        tot_sig += int(agg.get("valid_sig_count") or 0)
        tot_exec += int(agg.get("valid_exec_count") or 0)
        tot_exc += int(agg.get("excess_valid_count") or 0)
    # 分母 0 → coverage_by_basis 返回 None（契约 §8 约束 1），meta 写空单元格
    cov = (
        sc.coverage_by_basis(tot_sel, tot_sig, tot_exec, tot_exc)
        if have_agg
        else {sc.RET_BASIS_SIGNAL: None, sc.RET_BASIS_OPEN: None, "excess": None}
    )

    any_backfill = summary_backfill or week_backfill or detail_backfill
    notice = _backfill_notice() if any_backfill else None

    # 「指定规则补算」周（子集快照）逐周列明规则集：这些周的统计只有这几条
    # 规则，其余规则当周没有名单（不是当周空仓）。
    subset_weeks = [pw for pw in per_week if pw["scope"] == sc.RULES_SCOPE_SUBSET]
    subset_detail = "; ".join(
        f"{pw['wid']}: " + sc.subset_scope_notice(pw["scope_ids"])
        for pw in subset_weeks
    )

    if unsettled_weeks:
        warnings.append(f"{unsettled_weeks} 周暂无跟踪产物（未结算）")
    if subset_weeks:
        warnings.append(
            f"{len(subset_weeks)} 周为「指定规则补算」（仅部分规则）："
            "这些周其他规则没有名单与收益，跨周统计只反映各自被补算的规则"
        )

    # 写文件
    import openpyxl

    export_root = Path(cfg.storage_root) / EXPORT_SUBDIRNAME
    export_root.mkdir(parents=True, exist_ok=True)
    if entry_asof is not None:
        scope = f"week{int(entry_asof)}"
    elif rule_filter:
        scope = f"w{len(per_week)}_rule"
    else:
        scope = f"w{len(per_week)}_all"
    stamp = time.strftime("%Y%m%d_%H%M%S") + "_" + uuid.uuid4().hex[:6]
    filename = f"bagua_track_{scope}_{stamp}.xlsx"
    out = export_root / filename

    wb = openpyxl.Workbook()
    ws_sum = wb.active
    ws_sum.title = SHEET_SUMMARY
    ws_week = wb.create_sheet(SHEET_WEEK)
    ws_detail = wb.create_sheet(SHEET_DETAIL)

    _write_sheet(
        ws_sum,
        [
            "规则ID", "版本指纹", "跟踪周数", "已结算周数", "总票次",
            f"近{int(weeks)}周胜率(周等权,%)",  # = 各周 win_rate_sig 均值（同 L0 API weekly_equal_win_rate_sig）
            "平均收益(信号收盘,%)", "平均收益(首日开盘,%)",
            "平均超额(信号收盘,%)", "平均超额(首日开盘,%)",
            "有效周数", "样本是否充足",
        ],
        summary_out,
        notice if summary_backfill else None,
    )
    _write_sheet(
        ws_week,
        [
            "规则ID", "信号日(week_id)", "数据日(asof)", "运行来源(run_kind)",
            "入选数", "有效信号口径数", "有效开盘口径数", "待结算数",
            "缺数据数", "不可成交数", "未知数",
            "胜率(信号,%)", "平均收益(信号,%)", "平均收益(开盘,%)",
            "平均超额(信号收盘,%)", "平均超额(首日开盘,%)",
            "平均最大涨幅(%)", "平均回吐(%)", "完成状态(completion)", "是否回填",
        ],
        week_out,
        notice if week_backfill else None,
    )
    _write_sheet(
        ws_detail,
        [
            "规则ID", "信号日", "代码", "名称", "入场价(信号日收盘)", "首日开盘",
            "周五收盘", "可成交性",
            "周内最大涨幅(%)", "见顶日", "见顶星期", "最低相对入场收益(%)", "峰谷回撤(%)",
            "周五收益(信号口径,%)", "周五收益(开盘口径,%)", "理论开盘收益(%)",
            "基准收益(信号,%)", "基准收益(开盘,%)", "超额(信号,%)", "超额(开盘,%)",
            "状态",
        ],
        detail_out,
        notice if detail_backfill else None,
    )

    meta = wb.create_sheet(SHEET_META, 0)
    meta_rows: List[List[Any]] = [
        ["exported_at", time.strftime("%Y-%m-%d %H:%M:%S")],
        ["weeks", len(per_week)],
        ["rule_filter", rule_filter or ""],
        ["entry_asof", int(entry_asof) if entry_asof is not None else ""],
        ["snapshot_ids", ",".join(snapshot_ids)],
        ["tracking_revision_ids", ",".join(rev_ids)],
        ["data_version", _uniq_json(data_versions)],
        ["benchmark_data_version", _uniq_json(benchmark_versions)],
        ["algo_version", ",".join(algo_versions)],
        ["schema_version", ",".join(schema_versions)],
        ["coverage_signal_close", cov.get(sc.RET_BASIS_SIGNAL)],
        ["coverage_week_first_open", cov.get(sc.RET_BASIS_OPEN)],
        ["coverage_excess", cov.get("excess")],
        [
            "生成来源说明",
            "数据来源=发布索引指向的已发布快照 + 其当前 revision 的跟踪产物；"
            "统计口径见 docs/plans/auto-screen-track/contract.md §3；"
            "空仓周胜率/覆盖率与无有效样本一律为空单元格（null），不代表 0。",
        ],
    ]
    if notice:
        meta_rows.append(["backfill_notice", notice])
    if subset_weeks:
        # 子集周的规则范围逐周可追溯（读取方不必再回读快照）
        meta_rows.append(["subset_scope_weeks", "; ".join(str(pw["wid"]) for pw in subset_weeks)])
        meta_rows.append(["subset_scope_notice", subset_detail])
    meta.append([_excel_safe_cell("key"), _excel_safe_cell("value")])
    from openpyxl.styles import Font

    for cell in meta[1]:
        cell.font = Font(bold=True)
    for k, v in meta_rows:
        meta.append([_excel_safe_cell(k), _excel_safe_cell(v)])

    wb.save(out)

    rows_total = len(summary_out) + len(week_out) + len(detail_out)
    return {
        "ok": True,
        "file": filename,
        "path": str(out.resolve()),
        "download_url": "/api/v1/bagua/track/export/download?file="
        + filename,
        "sheets": list(SHEETS_ORDER),
        "rows_total": rows_total,
        "warnings": warnings,
    }
