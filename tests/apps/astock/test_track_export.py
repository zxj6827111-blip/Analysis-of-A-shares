# -*- coding: utf-8 -*-
"""阶段 3 测试：跟踪结果导出（service/track_export.py + 导出/下载端点）。

全部用合成快照 + 合成跟踪产物，不依赖真实数据根、不跑跟踪计算。
最小构造手法复用 tests/apps/astock/test_track_routes.py 的 _publish_snap/_write_track。
覆盖：正常导出、空值语义（null 非 0）、回填提示、meta 可追溯、无发布周、
下载端点安全白名单、参数边界。
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

import tests.apps.astock.conftest  # noqa: F401

from tests.apps.astock.test_track_routes import _publish_snap, _write_track
from wtpy.apps.astock.config import get_default_config
from wtpy.apps.astock.service.track_export import _weekday_cn

RID = "txt_测试规则A"


@pytest.fixture()
def track_client(tmp_path: Path):
    from fastapi.testclient import TestClient

    from wtpy.apps.astock.api import create_app

    storage = tmp_path / "st"
    ind = tmp_path / "ind"
    storage.mkdir(parents=True)
    ind.mkdir(parents=True)
    (ind / "测试规则A.txt").write_text("MA5:=MA(C,5);\nXG:CROSS(C,MA5);", encoding="utf-8")
    cfg = get_default_config(
        storage_root=storage, indicator_dir=ind, output_root=tmp_path / "out"
    )
    client = TestClient(create_app(cfg))
    yield client, cfg, storage
    client.close()


def _sheet_dicts(path: str, name: str):
    """按第一个表头单元格定位表头行，返回 (headers, [dict])。

    回填周的 sheet 顶部有提示行，不能假设表头在第 1 行。
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    assert name in wb.sheetnames, wb.sheetnames
    ws = wb[name]
    all_rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hidx = next(
        i for i, row in enumerate(all_rows) if row and row[0] == "规则ID"
    )
    headers = [("" if h is None else str(h)) for h in all_rows[hidx]]
    data = []
    for row in all_rows[hidx + 1:]:
        if not row or all(v is None for v in row):
            continue
        data.append(dict(zip(headers, row)))
    return headers, data


def _meta_dict(path: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["meta"]
    out = {}
    for row in ws.iter_rows(values_only=True):
        if row and row[0] not in (None, "key"):
            out[str(row[0])] = row[1] if len(row) > 1 else None
    return out


class TestExportNormal:
    def test_export_ok_sheets_and_percent(self, track_client):
        client, cfg, _ = track_client
        sid = _publish_snap(cfg, 20260911, ["SZSE.000001.SZ", "SZSE.000002.SZ"])
        _write_track(cfg, sid, 20260911)

        r = client.get("/api/v1/bagua/track/export")
        assert r.status_code == 200
        body = r.json()
        assert body["ok"] is True
        assert Path(body["path"]).exists()
        assert body["file"].startswith("bagua_track_")
        assert body["file"].endswith(".xlsx")
        assert set(body["sheets"]) == {"指标汇总", "周汇总", "周明细", "meta"}
        assert body["rows_total"] == 4  # 1 汇总 + 1 周汇总 + 2 明细

        wb = openpyxl.load_workbook(body["path"])
        assert set(wb.sheetnames) == {"指标汇总", "周汇总", "周明细", "meta"}

        _, summary = _sheet_dicts(body["path"], "指标汇总")
        assert len(summary) == 1
        srow = summary[0]
        assert srow["规则ID"] == RID
        assert srow["跟踪周数"] == 1 and srow["已结算周数"] == 1
        assert srow["总票次"] == 2
        # 百分数换算：mean_ret_close_sig=0.015 → 1.5
        assert srow["平均收益(信号收盘,%)"] == pytest.approx(1.5)
        assert srow["平均收益(首日开盘,%)"] == pytest.approx(3.0)
        assert "平均超额(信号收盘,%)" in srow
        assert "平均超额(首日开盘,%)" in srow
        # 周等权胜率 = 各周 win_rate_sig 的均值（与 L0 API 同口径；
        # 单周 win_rate_sig=0.5 → 50%——不是"正收益周占比"那种 100%）
        assert srow["近12周胜率(周等权,%)"] == pytest.approx(50.0)
        assert srow["样本是否充足"] == "样本不足"  # 2 < 30

        _, weeks = _sheet_dicts(body["path"], "周汇总")
        assert len(weeks) == 1
        wrow = weeks[0]
        assert wrow["信号日(week_id)"] == 20260911
        assert wrow["完成状态(completion)"] == "complete"
        assert wrow["入选数"] == 2
        assert wrow["是否回填"] == "否"
        assert "平均超额(信号收盘,%)" in wrow
        assert "平均超额(首日开盘,%)" in wrow
        assert wrow["平均最大涨幅(%)"] == pytest.approx(8.0)

        _, details = _sheet_dicts(body["path"], "周明细")
        assert len(details) == 2
        by_code = {d["代码"]: d for d in details}
        d1 = by_code["SZSE.000001.SZ"]
        # 0.05 → 5.00
        assert d1["周五收益(信号口径,%)"] == pytest.approx(5.0)
        assert d1["可成交性"] == "可成交"
        d2 = by_code["SZSE.000002.SZ"]
        assert d2["可成交性"] == "一字涨停买不进"
        assert d2["周五收益(开盘口径,%)"] is None  # 不可成交 → 留空

    def test_meta_traceable(self, track_client):
        client, cfg, _ = track_client
        sid = _publish_snap(cfg, 20260911, ["SZSE.000001.SZ"])
        _write_track(cfg, sid, 20260911, rev="rev1")
        body = client.get("/api/v1/bagua/track/export").json()
        meta = _meta_dict(body["path"])
        assert meta["snapshot_ids"] == sid
        assert meta["tracking_revision_ids"] == "rev1"
        assert meta["weeks"] == 1
        assert meta["schema_version"] == "1"
        assert meta["algo_version"] == "1"
        assert "生成来源说明" in meta

    def test_summary_win_rate_matches_l0_api(self, track_client):
        """★审查 🟡 对账：导出汇总胜率与 L0 API 的 weekly_equal_win_rate_sig
        必须是同一数字（同一列名两种算法 = UI 与 xlsx 对不上账）。"""
        client, cfg, _ = track_client
        # 两周同指纹段：胜率 0.5 / 1.0 → 周等权胜率 0.75
        sid1 = _publish_snap(cfg, 20260911, ["SZSE.000001.SZ", "SZSE.000002.SZ"])
        _write_track(cfg, sid1, 20260911, aggregates=[{
            "rule_id": RID, "selected_count": 2,
            "valid_sig_count": 2, "valid_exec_count": 2,
            "win_rate_sig": 0.5, "win_rate_exec": 0.5,
            "mean_ret_close_sig": 0.015, "mean_ret_close_exec": 0.02,
            "mean_excess_sig": None,
        }])
        sid2 = _publish_snap(cfg, 20260918, ["SZSE.000001.SZ"], fp="fpA")
        _write_track(cfg, sid2, 20260918, aggregates=[{
            "rule_id": RID, "selected_count": 1,
            "valid_sig_count": 1, "valid_exec_count": 1,
            "win_rate_sig": 1.0, "win_rate_exec": 1.0,
            "mean_ret_close_sig": 0.02, "mean_ret_close_exec": 0.02,
            "mean_excess_sig": None,
        }])
        body = client.get("/api/v1/bagua/track/export").json()
        _, summary = _sheet_dicts(body["path"], "指标汇总")
        l0 = client.get("/api/v1/bagua/track/rules").json()["rules"][0]
        assert len(summary) == 1
        assert summary[0]["近12周胜率(周等权,%)"] == pytest.approx(
            l0["weekly_equal_win_rate_sig"] * 100.0
        )
        assert summary[0]["平均收益(信号收盘,%)"] == pytest.approx(
            l0["weekly_equal_mean_ret_sig"] * 100.0
        )
        assert summary[0]["平均收益(首日开盘,%)"] == pytest.approx(
            l0["weekly_equal_mean_ret_exec"] * 100.0
        )


class TestEmptySemantics:
    def test_null_written_as_empty_not_zero(self, track_client):
        client, cfg, _ = track_client
        sid = _publish_snap(cfg, 20260911, [])
        _write_track(
            cfg, sid, 20260911, rows=[],
            aggregates=[{
                "rule_id": RID, "selected_count": 0,
                "valid_sig_count": 0, "valid_exec_count": 0,
                "win_rate_sig": None, "mean_ret_close_sig": None,
                "mean_ret_close_exec": None, "mean_excess_sig": None,
            }],
            coverage={"signal_close": None, "week_first_open": None, "excess": None},
        )
        body = client.get("/api/v1/bagua/track/export").json()
        assert body["ok"] is True

        _, summary = _sheet_dicts(body["path"], "指标汇总")
        srow = summary[0]
        assert srow["总票次"] == 0
        assert srow["有效周数"] == 0
        # 空仓周：胜率/均值是 null（无有效样本），不能写成 0
        assert srow["近12周胜率(周等权,%)"] is None
        assert srow["平均收益(信号收盘,%)"] is None
        assert srow["平均收益(首日开盘,%)"] is None
        assert srow["平均超额(信号收盘,%)"] is None
        assert srow["平均超额(首日开盘,%)"] is None

        _, weeks = _sheet_dicts(body["path"], "周汇总")
        assert weeks[0]["胜率(信号,%)"] is None
        assert weeks[0]["入选数"] == 0

        # 覆盖率分母为 0 → meta 写空单元格
        meta = _meta_dict(body["path"])
        assert meta["coverage_signal_close"] is None
        assert meta["coverage_week_first_open"] is None
        assert meta["coverage_excess"] is None


class TestBackfillNotice:
    def test_notice_on_backfill_sheets(self, track_client):
        client, cfg, _ = track_client
        sid = _publish_snap(cfg, 20260904, ["SZSE.000001.SZ"], run_kind="backfill")
        _write_track(cfg, sid, 20260904)
        body = client.get("/api/v1/bagua/track/export").json()
        assert body["ok"] is True

        wb = openpyxl.load_workbook(body["path"], data_only=True)
        for name in ("指标汇总", "周汇总", "周明细"):
            a1 = wb[name]["A1"].value
            assert a1 and "不代表当时实际发布名单" in a1, (name, a1)
        meta = _meta_dict(body["path"])
        assert "不代表当时实际发布名单" in str(meta["backfill_notice"])

        _, weeks = _sheet_dicts(body["path"], "周汇总")
        assert weeks[0]["是否回填"] == "是"


class TestNoPublishedWeek:
    def test_no_week_returns_reason_without_file(self, track_client):
        client, cfg, storage = track_client
        r = client.get("/api/v1/bagua/track/export")
        assert r.status_code == 200
        body = r.json()
        assert body == {"ok": False, "reason": "no_published_week"}
        # 不产空文件
        export_dir = storage / "bagua_exports"
        assert not export_dir.exists() or not list(export_dir.glob("*.xlsx"))


class TestDownloadEndpoint:
    def test_valid_download_200(self, track_client):
        client, cfg, _ = track_client
        sid = _publish_snap(cfg, 20260911, ["SZSE.000001.SZ"])
        _write_track(cfg, sid, 20260911)
        body = client.get("/api/v1/bagua/track/export").json()
        r = client.get(
            "/api/v1/bagua/track/export/download", params={"file": body["file"]}
        )
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers["content-type"]
        assert len(r.content) > 0

    @pytest.mark.parametrize(
        "bad",
        [
            "../../etc/passwd",
            "..%2f..%2fx.xlsx",
            "/etc/passwd",
            "C:\\Windows\\win.ini",
            "bagua_track_a.txt",       # 非 .xlsx
            "bagua_track_../x.xlsx",   # 含分隔符
        ],
    )
    def test_illegal_names_400(self, track_client, bad):
        client, _cfg, _ = track_client
        r = client.get("/api/v1/bagua/track/export/download", params={"file": bad})
        assert r.status_code == 400, bad

    def test_raw_percent_encoded_traversal_400(self, track_client):
        client, _cfg, _ = track_client
        r = client.get(
            "/api/v1/bagua/track/export/download?file=..%2f..%2fx.xlsx"
        )
        assert r.status_code == 400

    def test_missing_file_404(self, track_client):
        client, _cfg, _ = track_client
        r = client.get(
            "/api/v1/bagua/track/export/download",
            params={"file": "bagua_track_nope_000000.xlsx"},
        )
        assert r.status_code == 404


    def test_entry_asof_limits_single_week(self, track_client):
        """entry_asof 限定单周：两周都发布，只导出指定那一周。"""
        client, cfg, _ = track_client
        sid1 = _publish_snap(cfg, 20260911, ["SZSE.000001.SZ"])
        _write_track(cfg, sid1, 20260911)
        sid2 = _publish_snap(cfg, 20260918, ["SZSE.000002.SZ"], fp="fpB")
        _write_track(cfg, sid2, 20260918)
        body = client.get(
            "/api/v1/bagua/track/export", params={"entry_asof": "20260911"}
        ).json()
        assert body["ok"] is True
        assert "_week20260911_" in body["file"]
        _, weeks = _sheet_dicts(body["path"], "周汇总")
        assert [w["信号日(week_id)"] for w in weeks] == [20260911]


class TestParamValidation:
    @pytest.mark.parametrize("weeks", [0, 105])
    def test_weeks_out_of_range_400(self, track_client, weeks):
        client, _cfg, _ = track_client
        r = client.get("/api/v1/bagua/track/export", params={"weeks": weeks})
        assert r.status_code == 400

    def test_rule_id_too_long_400(self, track_client):
        client, _cfg, _ = track_client
        r = client.get(
            "/api/v1/bagua/track/export", params={"rule_id": "x" * 129}
        )
        assert r.status_code == 400

    def test_entry_asof_bad_format_400(self, track_client):
        client, _cfg, _ = track_client
        r = client.get(
            "/api/v1/bagua/track/export", params={"entry_asof": "2026"}
        )
        assert r.status_code == 400


class TestDetailColumnsV2:
    def test_detail_has_name_week_end_close_and_weekday(self, track_client):
        """schema v2 明细列：名称 / 周五收盘 / 见顶星期。

        UI 列表与导出消费同一份产物字段（名称、close_week_end、
        max_gain_sig_date），列必须一一对应——两端口径漂移会让用户
        对不上账。
        """
        client, cfg, _ = track_client
        sid = _publish_snap(cfg, 20260911, ["SZSE.000001.SZ"])
        _write_track(
            cfg, sid, 20260911,
            rows=[{
                "code": "SZSE.000001.SZ", "rule_id": RID, "name": "平安银行",
                "entry_close_signal": 10.0, "entry_open_week": 10.5,
                "close_week_end": 11.0, "fill_status": "ok",
                "ret_close_sig": 0.1, "ret_close_exec": 0.047619,
                "max_gain_sig": 0.08, "max_gain_sig_date": 20260916,
            }],
        )
        body = client.get("/api/v1/bagua/track/export").json()
        headers, detail = _sheet_dicts(body["path"], "周明细")
        for col in ("名称", "周五收盘", "见顶星期"):
            assert col in headers, headers
        row = detail[0]
        assert row["名称"] == "平安银行"
        assert row["周五收盘"] == pytest.approx(11.0)
        assert row["见顶星期"] == "周三"  # 2026-09-16 是周三
        # 缺日期的行：见顶星期留空串（不填占位、不猜）
        assert _weekday_cn(None) == ""
        assert _weekday_cn(0) == ""


# ---------------------------------------------------------------------------
# 导出与跟踪页同口径（2026-09-15 用户要求「与规则中心一致」）
# ---------------------------------------------------------------------------


class TestExportCatalogSync:
    """规则中心删掉的规则不导出；同一公式的多个 rule_id 归并成一行。

    导出与只读 API 必须同口径，否则 xlsx 的行数与页面显示对不上账。
    """

    @staticmethod
    def _patch_catalog(monkeypatch, ids):
        """注入规则目录判定；ids=None 表示目录不可得（不过滤不归并）。

        目录值与 api_routes.tracking._current_rule_catalog 同结构（含
        executable/hidden/source），导出与页面才能选出同一个 canonical。
        """
        from wtpy.apps.astock.service import track_export as te

        catalog = None
        if ids is not None:
            catalog = {
                rid: {"executable": True, "hidden": False, "source": "builtin"}
                for rid in ids
            }
        monkeypatch.setattr(te, "_current_rule_catalog", lambda _cfg: catalog)

    def test_export_skips_deleted_and_merges_same_fingerprint(
        self, track_client, monkeypatch
    ):
        from tests.apps.astock.test_track_routes import TestRuleCatalogSync

        client, cfg, _ = track_client
        sid = TestRuleCatalogSync._publish_multi(cfg, 20260911, [
            ("txt_已删除规则", "fpDel", ["SZSE.000001.SZ"]),
            ("txt_735金叉及趋势", "fpSame", ["SZSE.000002.SZ"]),
            ("tn6_735金叉及趋势", "fpSame", ["SZSE.000002.SZ"]),
        ])
        _write_track(cfg, sid, 20260911, rows=[
            {"code": "SZSE.000001.SZ", "rule_id": "txt_已删除规则",
             "ret_close_sig": 0.05, "ret_close_exec": 0.03, "fill_status": "ok"},
            {"code": "SZSE.000002.SZ", "rule_id": "txt_735金叉及趋势",
             "ret_close_sig": 0.02, "ret_close_exec": 0.01, "fill_status": "ok"},
            {"code": "SZSE.000002.SZ", "rule_id": "tn6_735金叉及趋势",
             "ret_close_sig": 0.02, "ret_close_exec": 0.01, "fill_status": "ok"},
        ])
        # 目录里只剩 tn6_（txt_ 被隐藏、已删规则不存在）
        self._patch_catalog(monkeypatch, {"tn6_735金叉及趋势"})
        body = client.get("/api/v1/bagua/track/export").json()
        assert body["ok"] is True
        _, summary = _sheet_dicts(body["path"], "指标汇总")
        assert [r["规则ID"] for r in summary] == ["tn6_735金叉及趋势"], \
            "已删规则不导出、同指纹归并成一行"
        headers, detail = _sheet_dicts(body["path"], "周明细")
        # 明细：已删规则的行不出现；同指纹组同周同票只留一行（规则ID=canonical）
        assert {r["规则ID"] for r in detail} == {"tn6_735金叉及趋势"}
        assert [r["代码"] for r in detail] == ["SZSE.000002.SZ"]
        assert body["rows_total"] == 1 + 1 + 1, "1 汇总 + 1 周汇总 + 1 明细"

    def test_export_keeps_all_when_catalog_unavailable(self, track_client, monkeypatch):
        """目录不可得（降级）→ 不过滤不归并（宁可多导，不能空表或报错）。"""
        from tests.apps.astock.test_track_routes import TestRuleCatalogSync

        client, cfg, _ = track_client
        sid = TestRuleCatalogSync._publish_multi(cfg, 20260911, [
            ("txt_A", "fpA", ["SZSE.000001.SZ"]),
            ("txt_B", "fpB", ["SZSE.000002.SZ"]),
        ])
        _write_track(cfg, sid, 20260911, rows=[
            {"code": "SZSE.000001.SZ", "rule_id": "txt_A",
             "ret_close_sig": 0.05, "fill_status": "ok"},
            {"code": "SZSE.000002.SZ", "rule_id": "txt_B",
             "ret_close_sig": -0.05, "fill_status": "ok"},
        ])
        self._patch_catalog(monkeypatch, None)
        body = client.get("/api/v1/bagua/track/export").json()
        _, summary = _sheet_dicts(body["path"], "指标汇总")
        assert sorted(r["规则ID"] for r in summary) == ["txt_A", "txt_B"]
