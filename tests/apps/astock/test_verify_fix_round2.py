# -*- coding: utf-8 -*-
"""独立复验（修复轮 2）：对 coder 的 6 项修复做独立端到端/桩验证。

1) indicator_review 显式默认规则 ID 的 sheet 名（735/5日外）
2) 空 stock_pool 跳过即时计算（不触发全市场扫描）
3) 规则 sheet 重名唯一化 + select_ids/missing_ids 去重
4) PUT queue/config 并发锁 + _parse_max_workers 整值 float
5) shutdown 后 set_max_workers 不扩容
6) bqExportReviewSuffix 在 review_note=="ok" 时不追加后缀

不依赖 coder 新用例的内部桩；不改业务代码。
"""
from __future__ import annotations

import datetime as _dt
import json
import re
import shutil
import subprocess
import threading
import time
from pathlib import Path
from unittest.mock import patch

import pytest

import tests.apps.astock.conftest  # noqa: F401
from tests.apps.astock.export_layout import data_rows
from tests.apps.astock.conftest import requires_real_formulas  # noqa: F401

from wtpy.apps.astock.config import AStockConfig, get_default_config
from wtpy.apps.astock.data.tdx_reader import DayBar
from wtpy.apps.astock.service.db import get_app_setting
from wtpy.apps.astock.service.jobs import JobStore

BACKEND_ROOT = Path(__file__).resolve().parents[3]
V3_HTML = (
    BACKEND_ROOT / "wtpy" / "apps" / "astock" / "web" / "static" / "index_v3.html"
)
BAGUA_JSON = (
    BACKEND_ROOT / "wtpy" / "apps" / "astock" / "bagua" / "bagua_384.json"
)

_DS_META_MOCK = {
    "dataset_id": "mock",
    "dataset_source": "tdxquant",
    "dataset_adjustment": "front",
    "dataset_status": "ready",
    "covers_asof": True,
    "candidate_datasets": 1,
}


def _weekdays(start: str, end: str) -> list:
    out = []
    cur = _dt.date.fromisoformat(start)
    end_d = _dt.date.fromisoformat(end)
    while cur <= end_d:
        if cur.isoweekday() <= 5:
            out.append(int(cur.strftime("%Y%m%d")))
        cur += _dt.timedelta(days=1)
    return out


def _make_app(cfg: AStockConfig):
    from wtpy.apps.astock.api import create_app

    return create_app(cfg)


def _make_client(app):
    from fastapi.testclient import TestClient

    return TestClient(app)


@pytest.fixture()
def api_cfg(tmp_path: Path) -> AStockConfig:
    storage = tmp_path / "st"
    ind = tmp_path / "ind"
    storage.mkdir()
    ind.mkdir()
    return get_default_config(
        storage_root=storage, indicator_dir=ind, output_root=tmp_path / "out"
    )


# ===========================================================================
# 修复 1：显式默认规则 ID 的 sheet 名
# ===========================================================================


@requires_real_formulas
def test_round2_explicit_default_rule_ids_sheet_names_reversed_order(tmp_path):
    """显式传默认两条（且与默认顺序相反）：sheet 仍精确为 735 / 5日外，
    不得退化为完整 rule_id；顺序按传入顺序保留。"""
    from wtpy.apps.astock.service import indicator_review as ir

    from tests.apps.astock.conftest import formula_cfg

    cfg = formula_cfg(tmp_path)
    days = _weekdays("2026-06-01", "2026-08-28")
    bars = [DayBar(d, 10.0, 10.1, 9.9, 10.0, 1e6, 1e7) for d in days]

    def _loader(code, asof):
        return [b for b in bars if int(b.date) <= int(asof)], {"dataset_id": "mock"}

    def _surface(_cfg):
        return {"formal_l1_id": "mock_l1", "max_date": 20260828}, ""

    out = ir.run_weekly_review(
        cfg,
        asof=20260828,
        codes=["SSE.STK.600000"],
        rule_ids=["txt_先跌后涨新版5日外", "txt_735金叉及趋势"],
        persist=True,
        bar_loader=_loader,
        surface_resolver=_surface,
    )
    assert out["status"] == "ok"
    pairs = [(r["rule_id"], r["sheet"]) for r in out["rules"]]
    assert pairs == [
        ("txt_先跌后涨新版5日外", "5日外"),
        ("txt_735金叉及趋势", "735"),
    ]
    # 反向断言：绝不能是完整 rule_id（修复前行为）
    assert all(sheet in ("735", "5日外") for _rid, sheet in pairs)
    # 落盘 JSON（周五链共享产物）同样必须是短名
    on_disk = json.loads(
        ir.review_output_path(cfg, 20260828).read_text(encoding="utf-8")
    )
    assert [(r["rule_id"], r["sheet"]) for r in on_disk["rules"]] == pairs


# ===========================================================================
# 修复 2：空 stock_pool 跳过即时计算
# ===========================================================================


def _write_review(cfg: AStockConfig, asof: int, payload: dict) -> None:
    d = Path(cfg.storage_root) / "indicator_review"
    d.mkdir(parents=True, exist_ok=True)
    (d / f"review_{asof}.json").write_text(
        json.dumps(payload, ensure_ascii=False), encoding="utf-8"
    )


def _review_payload(asof: int, rules: list) -> dict:
    return {
        "asof": int(asof),
        "generated_at": "2026-08-28 19:00:00",
        "status": "ok",
        "no_go_reason": "",
        "universe_size": 1,
        "scanned": 1,
        "error_count": 0,
        "rules": rules,
    }


def _mock_etf_export(monkeypatch, cfg: AStockConfig):
    """ETF-only 导出桩：空股票池 + 单只 ETF 有行情。"""
    from wtpy.apps.astock.service import bagua_query as bq

    etf_days = _weekdays("2026-06-01", "2026-08-28")
    etf_bars = [DayBar(d, 4.05, 4.12, 3.98, 4.06, 1.0, 1.0) for d in etf_days]
    monkeypatch.setattr(
        bq,
        "_load_dataset_bars",
        lambda *_a, **_k: (etf_bars, dict(_DS_META_MOCK)),
    )
    monkeypatch.setattr(
        bq,
        "BaguaPlaneSession",
        lambda *_a, **_k: (_ for _ in ()).throw(FileNotFoundError("no md")),
    )
    monkeypatch.setattr(bq, "list_etf_std_codes", lambda _cfg: [])
    return cfg


def test_round2_empty_stock_pool_skips_compute_and_full_market(tmp_path, monkeypatch):
    """stock_pool 为空 + 勾选未预计算规则：`_compute_rules_for_export` 与
    `run_weekly_review`（全市场扫描入口）都不得被调用；导出成功，note 有 skip。"""
    if not BAGUA_JSON.exists():
        pytest.skip("bagua_384.json missing")

    cfg = get_default_config(
        storage_root=tmp_path / "st",
        indicator_dir=tmp_path / "ind",
        output_root=tmp_path / "out",
    )
    Path(cfg.storage_root).mkdir(parents=True, exist_ok=True)
    Path(cfg.indicator_dir).mkdir(parents=True, exist_ok=True)
    _mock_etf_export(monkeypatch, cfg)

    # 复核文件只有 5日外，勾选 735 → 735 属于 missing（未预计算）
    _write_review(
        cfg,
        20240115,
        _review_payload(
            20240115,
            [
                {
                    "rule_id": "txt_先跌后涨新版5日外",
                    "sheet": "5日外",
                    "count": 0,
                    "matched": [],
                }
            ],
        ),
    )

    calls: list = []

    def _compute_stub(cfg_, asof, rule_ids, *, codes=None, on_progress=None):
        calls.append(("compute", list(rule_ids), list(codes or [])))
        raise AssertionError("空 stock_pool 不得触发即时计算/全市场扫描")

    from wtpy.apps.astock.service import bagua_query as bq
    from wtpy.apps.astock.service import indicator_review as ir

    monkeypatch.setattr(bq, "_compute_rules_for_export", _compute_stub)
    monkeypatch.setattr(
        ir,
        "run_weekly_review",
        lambda *_a, **_k: (_ for _ in ()).throw(
            AssertionError("空 stock_pool 不得进入全市场扫描")
        ),
    )

    info: dict = {}
    path = bq.export_bagua_multi_period_xlsx(
        cfg,
        date="2024-01-15",
        periods=["WEEK", "MONTH"],
        adjust="tushare_qfq",
        codes=["sh510300"],  # ETF-only → stock_pool == []
        all_stocks=False,
        review_rules=["txt_735金叉及趋势"],
        info_out=info,
    )

    assert calls == []
    assert path.exists()

    import openpyxl

    wb = openpyxl.load_workbook(path)
    assert set(wb.sheetnames) == {"meta", "index-all", "etf-all"}, wb.sheetnames
    meta = {r[0]: r[1] for r in wb["meta"].iter_rows(min_row=2, values_only=True)}
    note = str(meta["indicator_review_note"])
    assert "skip:导出票池为空，信号规则未计算" in note
    assert meta["indicator_review_sheets"] in ("", None)
    assert "skip:导出票池为空" in str(info["review_note"])
    assert info["query_date"] == 20240115


def test_round2_empty_stock_pool_never_enters_run_weekly_review(
    tmp_path, monkeypatch
):
    """更深的可观测点：不替换 `_compute_rules_for_export`，只监听其内部
    真正的全市场入口 `indicator_review.run_weekly_review`——必须一次都不进。"""
    if not BAGUA_JSON.exists():
        pytest.skip("bagua_384.json missing")

    cfg = get_default_config(
        storage_root=tmp_path / "st",
        indicator_dir=tmp_path / "ind",
        output_root=tmp_path / "out",
    )
    Path(cfg.storage_root).mkdir(parents=True, exist_ok=True)
    Path(cfg.indicator_dir).mkdir(parents=True, exist_ok=True)
    _mock_etf_export(monkeypatch, cfg)
    _write_review(cfg, 20240115, _review_payload(20240115, []))

    entered: list = []

    def _scan_entry(*args, **kwargs):
        entered.append((args, kwargs))
        raise AssertionError("ETF-only 导出不得进入全市场复核扫描")

    from wtpy.apps.astock.service import bagua_query as bq
    from wtpy.apps.astock.service import indicator_review as ir

    monkeypatch.setattr(ir, "run_weekly_review", _scan_entry)

    info: dict = {}
    path = bq.export_bagua_multi_period_xlsx(
        cfg,
        date="2024-01-15",
        periods=["WEEK", "MONTH"],
        adjust="tushare_qfq",
        codes=["sh510300"],
        all_stocks=False,
        review_rules=["txt_735金叉及趋势"],
        info_out=info,
    )
    assert entered == [], "run_weekly_review 被调用（会触发全市场扫描）"
    assert path.exists()
    assert "skip:导出票池为空" in str(info["review_note"])


# ===========================================================================
# 修复 3：规则 sheet 重名唯一化 + 选择去重
# ===========================================================================

def _mock_stock_export(monkeypatch, tmp_path):
    """全市场导出桩：两只股票，无 ETF，卦象面走 _load_dataset_bars。"""
    from wtpy.apps.astock.service import bagua_query as bq

    days = _weekdays("2023-12-01", "2024-01-15")
    bars = [DayBar(d, 6.27, 7.33, 5.90, 5.90, 1.0, 1.0) for d in days]
    monkeypatch.setattr(
        bq, "_load_dataset_bars", lambda *_a, **_k: (bars, dict(_DS_META_MOCK))
    )
    monkeypatch.setattr(
        bq,
        "BaguaPlaneSession",
        lambda *_a, **_k: (_ for _ in ()).throw(FileNotFoundError("no md")),
    )
    monkeypatch.setattr(
        bq,
        "_resolve_batch_codes",
        lambda cfg, codes=None, *, all_stocks=False: [
            "SSE.STK.600000",
            "SSE.STK.000001",
        ],
    )
    monkeypatch.setattr(bq, "list_etf_std_codes", lambda cfg: [])
    return tmp_path


def test_round2_duplicate_sheet_names_from_compute_and_select_dedup(
    tmp_path, monkeypatch
):
    """勾选列表重复 + 即时计算返回同名 sheet：只出 2 张互不相同的 sheet，
    成员各归其主，计算入参已完成去重。"""
    if not BAGUA_JSON.exists():
        pytest.skip("bagua_384.json missing")

    cfg = get_default_config(
        storage_root=tmp_path / "st",
        indicator_dir=tmp_path / "ind",
        output_root=tmp_path / "out",
    )
    Path(cfg.storage_root).mkdir(parents=True, exist_ok=True)
    Path(cfg.indicator_dir).mkdir(parents=True, exist_ok=True)
    _mock_stock_export(monkeypatch, tmp_path)

    # 复核文件：ok 但没有任何命中规则 → 勾选的 user_a/user_b 全走即时计算
    _write_review(cfg, 20240115, _review_payload(20240115, []))

    called: dict = {}

    def _compute_stub(cfg_, asof, rule_ids, *, codes=None, on_progress=None):
        called["rule_ids"] = list(rule_ids)
        return {
            "asof": asof,
            "status": "ok",
            "rules": [
                {
                    "rule_id": "user_a",
                    "sheet": "同名规则",
                    "count": 1,
                    "matched": [{"code": "SSE.STK.600000", "close": 5.9}],
                },
                {
                    "rule_id": "user_b",
                    "sheet": "同名规则",
                    "count": 1,
                    "matched": [{"code": "SSE.STK.000001", "close": 5.9}],
                },
            ],
        }

    from wtpy.apps.astock.service import bagua_query as bq

    monkeypatch.setattr(bq, "_compute_rules_for_export", _compute_stub)

    path = bq.export_bagua_multi_period_xlsx(
        cfg,
        date="2024-01-15",
        periods=["WEEK", "MONTH"],
        adjust="tushare_qfq",
        all_stocks=True,
        # 同一规则重复传两次 + 第二条
        review_rules=["user_a", "user_a", "user_b"],
    )

    # 选择/缺失列表去重：计算只收到一次 user_a、一次 user_b
    assert called["rule_ids"] == ["user_a", "user_b"]

    import openpyxl

    wb = openpyxl.load_workbook(path)
    signal_sheets = [
        n for n in wb.sheetnames if n not in ("meta", "stock-all", "index-all")
    ]
    assert len(signal_sheets) == 2, wb.sheetnames
    assert len(set(signal_sheets)) == 2
    assert all(len(n) <= 31 for n in signal_sheets)
    rows = {n: [r[0] for r in data_rows(wb[n])] for n in signal_sheets}
    by_members = {tuple(v): k for k, v in rows.items()}
    assert ("600000",) in by_members and ("000001",) in by_members
    assert by_members[("600000",)] != by_members[("000001",)]
    assert "同名规则" in by_members[("600000",)]


def test_round2_duplicate_selection_precomputed_path_single_sheet(
    tmp_path, monkeypatch
):
    """预计算路径：同一规则在 review_rules 里重复两次，只出 1 张 sheet，
    meta.indicator_review_rules_selected 也去重。"""
    if not BAGUA_JSON.exists():
        pytest.skip("bagua_384.json missing")

    cfg = get_default_config(
        storage_root=tmp_path / "st",
        indicator_dir=tmp_path / "ind",
        output_root=tmp_path / "out",
    )
    Path(cfg.storage_root).mkdir(parents=True, exist_ok=True)
    Path(cfg.indicator_dir).mkdir(parents=True, exist_ok=True)
    _mock_stock_export(monkeypatch, tmp_path)
    _write_review(
        cfg,
        20240115,
        _review_payload(
            20240115,
            [
                {
                    "rule_id": "txt_735金叉及趋势",
                    "sheet": "735",
                    "count": 1,
                    "matched": [{"code": "SSE.STK.600000", "close": 5.9}],
                },
                {
                    "rule_id": "txt_先跌后涨新版5日外",
                    "sheet": "5日外",
                    "count": 0,
                    "matched": [],
                },
            ],
        ),
    )

    from wtpy.apps.astock.service import bagua_query as bq

    path = bq.export_bagua_multi_period_xlsx(
        cfg,
        date="2024-01-15",
        periods=["WEEK", "MONTH"],
        adjust="tushare_qfq",
        all_stocks=True,
        review_rules=["txt_735金叉及趋势", "txt_735金叉及趋势"],
    )

    import openpyxl

    wb = openpyxl.load_workbook(path)
    assert "735" in wb.sheetnames
    assert [n for n in wb.sheetnames if n.startswith("735~")] == []
    signal_sheets = [
        n for n in wb.sheetnames if n not in ("meta", "stock-all", "index-all")
    ]
    assert signal_sheets == ["735"], wb.sheetnames
    meta = {r[0]: r[1] for r in wb["meta"].iter_rows(min_row=2, values_only=True)}
    assert meta["indicator_review_rules_selected"] == "txt_735金叉及趋势"


# ===========================================================================
# 修复 4：PUT 并发锁 + _parse_max_workers 整值 float
# ===========================================================================

def test_round2_parse_max_workers_route_integral_float_and_rejects(
    api_cfg: AStockConfig, monkeypatch
):
    """路由级：3.0 接受；2.5/True 400；_parse_max_workers docstring 语义一致。"""
    pytest.importorskip("fastapi")
    monkeypatch.delenv("ASTOCK_BT_MAX_WORKERS", raising=False)
    from wtpy.apps.astock.api_routes.backtests import _parse_max_workers

    assert _parse_max_workers(3.0) == 3
    assert _parse_max_workers(" 4 ") == 4
    for bad in (True, 2.5, "abc", None, {}):
        with pytest.raises((ValueError, TypeError)):
            _parse_max_workers(bad)

    app = _make_app(api_cfg)
    client = _make_client(app)
    jobs = app.state.astock.jobs
    try:
        r = client.put("/api/v1/backtests/queue/config", json={"max_workers": 3.0})
        assert r.status_code == 200, r.text
        assert r.json()["max_workers"] == 3
        assert jobs.max_workers == 3
        for bad in (2.5, True):
            rb = client.put(
                "/api/v1/backtests/queue/config", json={"max_workers": bad}
            )
            assert rb.status_code == 400, (bad, rb.status_code)
        assert jobs.max_workers == 3
    finally:
        jobs.shutdown(wait=True)


def test_round2_concurrent_put_db_and_runtime_stay_consistent(
    api_cfg: AStockConfig, monkeypatch
):
    """两线程并发 PUT 3 与 5 多轮：最终 DB 值 ∈ {3,5}，且
    ctx.jobs.max_workers 与 GET config 必须与 DB 完全一致（锁保证）。"""
    pytest.importorskip("fastapi")
    monkeypatch.delenv("ASTOCK_BT_MAX_WORKERS", raising=False)
    app = _make_app(api_cfg)
    ctx = app.state.astock
    client_a = _make_client(app)
    client_b = _make_client(app)
    try:
        for rnd in range(8):
            barrier = threading.Barrier(2)
            results: dict = {}

            def _put(client, n):
                barrier.wait(timeout=5.0)
                results[n] = client.put(
                    "/api/v1/backtests/queue/config", json={"max_workers": n}
                )

            t1 = threading.Thread(target=_put, args=(client_a, 3))
            t2 = threading.Thread(target=_put, args=(client_b, 5))
            t1.start()
            t2.start()
            t1.join(10.0)
            t2.join(10.0)
            assert not t1.is_alive() and not t2.is_alive()
            assert set(results) == {3, 5}, rnd
            for n, resp in results.items():
                assert resp.status_code == 200, (rnd, n, resp.text)

            db_val = get_app_setting(api_cfg, "bt_max_workers")
            assert db_val in ("3", "5"), db_val
            assert ctx.jobs.max_workers == int(db_val), (rnd, db_val, ctx.jobs.max_workers)

            d = client_a.get("/api/v1/backtests/queue/config").json()
            assert d["source"] == "db"
            assert d["max_workers"] == int(db_val)
    finally:
        ctx.jobs.shutdown(wait=True)


def test_round2_concurrent_put_serialized_by_lock_with_slow_io(
    api_cfg: AStockConfig, monkeypatch
):
    """放大竞态窗口的确定性验证：让 DB 写入按值快慢不同（3 慢 5 快），
    若无锁则会出现「DB=5 而运行时=3」的交错；持锁后每轮结束必须一致。"""
    pytest.importorskip("fastapi")
    monkeypatch.delenv("ASTOCK_BT_MAX_WORKERS", raising=False)
    app = _make_app(api_cfg)
    ctx = app.state.astock
    client_a = _make_client(app)
    client_b = _make_client(app)

    from wtpy.apps.astock.api_routes import backtests as bt_routes

    orig_set = bt_routes.set_app_setting

    def _slow_set(cfg, key, value):
        orig_set(cfg, key, value)
        # 写 3 的请求睡更久，迫使「后写库者先改运行时」的交错
        time.sleep(0.3 if str(value) == "3" else 0.05)

    orig_resize = JobStore.set_max_workers

    def _slow_resize(self, n):
        time.sleep(0.05)
        return orig_resize(self, n)

    monkeypatch.setattr(bt_routes, "set_app_setting", _slow_set)
    monkeypatch.setattr(JobStore, "set_max_workers", _slow_resize)

    try:
        for rnd in range(3):
            barrier = threading.Barrier(2)
            results: dict = {}

            def _put(client, n):
                barrier.wait(timeout=5.0)
                results[n] = client.put(
                    "/api/v1/backtests/queue/config", json={"max_workers": n}
                )

            t1 = threading.Thread(target=_put, args=(client_a, 3))
            t2 = threading.Thread(target=_put, args=(client_b, 5))
            t1.start()
            t2.start()
            t1.join(15.0)
            t2.join(15.0)
            assert not t1.is_alive() and not t2.is_alive()
            assert set(results) == {3, 5}

            db_val = get_app_setting(api_cfg, "bt_max_workers")
            assert db_val in ("3", "5"), (rnd, db_val)
            assert ctx.jobs.max_workers == int(db_val), (
                f"round {rnd}: DB={db_val} runtime={ctx.jobs.max_workers} "
                "（无锁交错的典型症状）"
            )
    finally:
        ctx.jobs.shutdown(wait=True)


# ===========================================================================
# 修复 5：shutdown 后 set_max_workers 不扩容
# ===========================================================================


def test_round2_set_max_workers_after_shutdown_returns_current_no_growth(
    tmp_path: Path,
):
    cfg = AStockConfig()
    cfg.storage_root = tmp_path / "store"
    cfg.output_root = tmp_path / "out"
    cfg.storage_root.mkdir(parents=True)
    cfg.output_root.mkdir(parents=True)

    store = JobStore(cfg, max_workers=2)
    before_threads = list(store._workers)
    assert len(before_threads) == 2
    store.shutdown(wait=True)
    assert not [t for t in store._workers if t.is_alive()]

    # shutdown 后任意扩容请求：返回值/上限不变、线程对象不新增、无存活线程
    assert store.set_max_workers(4) == 2
    assert store.max_workers == 2
    assert store._workers == before_threads, "不得再生产新线程对象"
    assert not [t for t in store._workers if t.is_alive()]
    # 再次调用同样无副作用
    assert store.set_max_workers(8) == 2
    assert not [t for t in store._workers if t.is_alive()]


# ===========================================================================
# 修复 6：前端 note=="ok" 不追加后缀（功能级 Node 执行）
# ===========================================================================


def _extract_js_function(src: str, name: str) -> str:
    m = re.search(r"(?:async\s+)?function\s+" + name + r"\s*\(", src)
    assert m, f"function not found: {name}"
    brace = src.index("{", m.start())
    depth = 0
    for j in range(brace, len(src)):
        if src[j] == "{":
            depth += 1
        elif src[j] == "}":
            depth -= 1
            if depth == 0:
                return src[m.start(): j + 1]
    raise AssertionError(f"unbalanced {name}")


def test_round2_bq_export_review_suffix_ok_is_empty_node():
    """真跑 JS：note="ok" + 非回退 → 空后缀；回退/异常 note 仍产出后缀。"""
    node = shutil.which("node")
    assert node, "node required"
    src = V3_HTML.read_text(encoding="utf-8")
    fn = _extract_js_function(src, "bqExportReviewSuffix")
    script = (
        fn
        + "\n"
        + """
function assert(cond, msg) { if (!cond) throw new Error(msg); }
// note=ok 且非回退：不得出现（ok）后缀
assert(bqExportReviewSuffix({ review_note: "ok", review_asof_used: 20240115, query_date: 20240115 }) === "", "ok must yield empty suffix");
// note=ok 但确有回退：用日期合成文案，而不是（ok）
var s1 = bqExportReviewSuffix({ review_note: "ok", review_fallback: 1, review_asof_used: 20240112, query_date: 20240115 });
assert(s1.indexOf("ok") < 0 && s1.indexOf("20240112") >= 0, "fallback must synthesize, got: " + s1);
// 真实回退 note：原样加括号
var s2 = bqExportReviewSuffix({ review_note: "fallback:使用 20240112 复核", review_asof_used: 20240112, query_date: 20240115 });
assert(s2 === "（fallback:使用 20240112 复核）", "note passthrough got: " + s2);
// skip 说明也应展示
var s3 = bqExportReviewSuffix({ review_note: "ok；skip:导出票池为空，信号规则未计算" });
assert(s3.indexOf("skip:") >= 0, "skip note must show, got: " + s3);
console.log("PASS bqExportReviewSuffix");
"""
    )
    proc = subprocess.run(
        [node, "-e", script],
        capture_output=True,
        text=True,
        cwd=str(BACKEND_ROOT),
    )
    assert proc.returncode == 0, proc.stdout + "\n" + proc.stderr
    assert "PASS bqExportReviewSuffix" in proc.stdout
