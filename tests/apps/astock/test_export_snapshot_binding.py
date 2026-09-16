# -*- coding: utf-8 -*-
"""阶段 1e：批量导出的「预计算规则 sheet」读不可变快照（双路径分离）。

验收口径（docs/plans/auto-screen-track/contract.md §0 + 评审第 7 条）：
- review_rules=None（旧调用方）一行不改，继续读 review_{asof}.json 兼容层；
- 显式勾选规则的导出优先读**已发布快照**（published 指针）：
  * 全部 covered → 信号 sheet 全部 source=snapshot，meta 记 snapshot_id；
  * 部分 covered → covered 用快照、其余即时计算（不失败、不静默丢 sheet）；
  * 无快照/快照不健康（status 非 ok、week 归属不符）→ 当作无快照走旧路径；
- 一致性：同 snapshot_id + 同规则集 + 同票池下，导出命中 == 筛选接口
  （combine_snapshot_hits）的 hits——两者读同一快照文件做同样的票池交集。

测试层级说明：优先走真实导出服务函数 export_bagua_multi_period_xlsx
（mock 日线/票池/注册表，与 test_bagua_query.py 的 _export_review_monkeypatch
同口径），断言落在生成的 xlsx 上；快照用 write_and_publish_snapshot 真实
发布（O_EXCL + 索引指针），指纹取自当前注册表（真实陈旧判定路径）。
纯函数（_snapshot_rules_to_review_form / _load_export_snapshot_for_week）
再各给一条单测锁语义。
"""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest

import tests.apps.astock.conftest  # noqa: F401
from tests.apps.astock.export_layout import data_rows

from wtpy.apps.astock.bagua.calculator import BaguaCalculator
from wtpy.apps.astock.data.tdx_reader import DayBar
from wtpy.apps.astock.service import bagua_query as bq
from wtpy.apps.astock.service import screen_contract as sc
from wtpy.apps.astock.service import screen_snapshots as ss

JSON_PATH = (
    Path(__file__).resolve().parents[3]
    / "wtpy"
    / "apps"
    / "astock"
    / "bagua"
    / "bagua_384.json"
)

_DS_META_MOCK = {
    "dataset_id": "mock",
    "dataset_source": "tdxquant",
    "dataset_adjustment": "front",
    "dataset_status": "ready",
    "covers_asof": True,
    "candidate_datasets": 1,
}

# 测试 universe：与 test_screen_cache_first 同三只票（SZSE 6 位展示码 000001…）
UNIVERSE = ["SZSE.000001.SZ", "SZSE.000002.SZ", "SZSE.000003.SZ"]
ASOF = 20260911  # 周五（信号日 = 周内最后交易日）


# ---------------------------------------------------------------------------
# 脚手架（与 test_bagua_query.py::_export_review_monkeypatch 同口径）
# ---------------------------------------------------------------------------


def _export_env(monkeypatch, tmp_path):
    """最小导出环境：两只股票、无 ETF、卦象面走 _load_dataset_bars。

    不落 review_{asof}.json（默认走快照/现算路径时的「无复核文件」基线，
    与回归用例的语义差异：那些用例靠 review JSON 提供预计算，本文件靠快照）。
    """
    if not JSON_PATH.exists():
        pytest.skip("bagua_384.json missing")

    stock_bars = [DayBar(ASOF, 6.27, 7.33, 5.90, 5.90, 1.0, 1.0)]

    def _fake_load(cfg, std_code, source_key, asof=None, **_kw):
        return stock_bars, dict(_DS_META_MOCK)

    monkeypatch.setattr(bq, "_load_dataset_bars", _fake_load)
    monkeypatch.setattr(
        bq,
        "BaguaPlaneSession",
        lambda *_a, **_k: (_ for _ in ()).throw(FileNotFoundError("no md")),
    )
    monkeypatch.setattr(
        bq,
        "_resolve_batch_codes",
        lambda cfg, codes=None, *, all_stocks=False: [
            "SZSE.000001.SZ",
            "SZSE.000002.SZ",
        ],
    )
    monkeypatch.setattr(bq, "list_etf_std_codes", lambda cfg: [])

    # 公式目录 = tmp 副本（isolated_indicator_dir）：能解析默认复核公式，
    # 指纹走真实注册表，又不污染真实 指标/ 目录
    from tests.apps.astock.conftest import isolated_indicator_dir

    ind_dir = isolated_indicator_dir(tmp_path)
    cfg = SimpleNamespace(
        bagua_json=JSON_PATH,
        storage_root=tmp_path,
        tdx_root=tmp_path,
        market_data_root=tmp_path / "md",
        forecast_root=tmp_path,
        forecast_weekly_dir=tmp_path,
        universe_path=tmp_path / "universe.json",
        adj_root=tmp_path,
        indicator_dir=ind_dir,
        mapping_path=Path(ind_dir) / "_no_tn6_map.json",
        calendar_path=tmp_path / "calendar.json",
    )
    return cfg


def _publish_snapshot(
    cfg,
    *,
    rules=None,
    asof=ASOF,
    status="ok",
    week_id=None,
):
    """合成并真实发布一个快照（write_and_publish_snapshot：O_EXCL + 指针）。

    规则指纹必须用当前注册表的真实指纹——导出/筛选的覆盖判定都比对
    当前公式指纹，合成指纹会被判 stale 落回现算（契约的防陈旧语义）。
    rules: [(rule_id, sheet, hit_codes)]，默认两条预置规则（735/5日外）。
    """
    from wtpy.apps.astock.service.screening import _current_rule_fingerprints

    if rules is None:
        rules = [
            ("txt_735金叉及趋势", "735", ("SZSE.000001.SZ",)),
            ("txt_先跌后涨新版5日外", "5日外", ("SZSE.000002.SZ",)),
        ]
    rid_list = [r[0] for r in rules]
    fps = _current_rule_fingerprints(cfg, rid_list) or {r: "fp" for r in rid_list}
    snap_rules = [
        {
            "rule_id": rid,
            "sheet": sheet,
            "status": "ok",
            "count": len(hits),
            "matched": [{"code": c, "close": 10.0} for c in hits],
            "failed_codes": [],
        }
        for rid, sheet, hits in rules
    ]
    payload = {
        "schema_version": ss.SNAPSHOT_SCHEMA_VERSION,
        "snapshot_id": sc.new_snapshot_id(asof),
        "run_kind": "weekly_chain",
        "week_id": asof if week_id is None else week_id,
        "asof": asof,
        "generated_at": "2026-09-11 18:40:00",
        "status": status,
        "universe_size": len(UNIVERSE),
        "universe_codes": list(UNIVERSE),
        "universe_fingerprint": "ufp1",
        "name_snapshot_id": "ns1",
        "rule_fingerprints": fps,
        "data_version": {},
        "content_fingerprint": "cfp1",
        "scanned": len(UNIVERSE),
        "missing_count": 0,
        "no_data_codes": [],
        "rules": snap_rules,
        "duration_sec": 0.1,
    }
    return ss.write_and_publish_snapshot(cfg, payload)


def _meta_of(wb):
    return {r[0]: r[1] for r in wb["meta"].iter_rows(min_row=2, values_only=True)}


def _load_wb(path):
    import openpyxl

    return openpyxl.load_workbook(path)


# ---------------------------------------------------------------------------
# 1. 快照覆盖全部请求规则 → 全部 source=snapshot
# ---------------------------------------------------------------------------


def test_full_coverage_all_sheets_from_snapshot(monkeypatch, tmp_path):
    """勾选的两条规则快照全覆盖：零即时计算、零 review JSON，
    meta 有 snapshot_id，rule_sources 逐条 snapshot:{asof}。"""
    cfg = _export_env(monkeypatch, tmp_path)

    def _no_compute(*_a, **_k):
        raise AssertionError("快照全覆盖时不得触发即时计算")

    monkeypatch.setattr(bq, "_compute_rules_for_export", _no_compute)
    from wtpy.apps.astock.service import indicator_review as ir

    def _no_review_load(*_a, **_k):
        raise AssertionError("快照路径命中时不得读 review JSON")

    monkeypatch.setattr(ir, "load_review_for_export", _no_review_load)

    pub = _publish_snapshot(cfg)
    assert pub["published"] is True
    snap_id = pub["snapshot_id"]

    path = bq.export_bagua_multi_period_xlsx(
        cfg, date=ASOF, periods=["WEEK", "MONTH"],
        adjust="tushare_qfq", all_stocks=True,
        review_rules=["txt_735金叉及趋势", "txt_先跌后涨新版5日外"],
    )
    wb = _load_wb(path)
    assert "735" in wb.sheetnames and "5日外" in wb.sheetnames
    # 命中 ∩ 票池（票池=000001/000002）：735 命中 000001，5日外命中 000002
    assert [r[0] for r in data_rows(wb["735"])] == ["000001"]
    assert [r[0] for r in data_rows(wb["5日外"])] == ["000002"]
    meta = _meta_of(wb)
    assert meta["indicator_review_snapshot_id"] == snap_id
    sources = str(meta["indicator_review_rule_sources"])
    assert sources == (
        "txt_735金叉及趋势=snapshot:20260911;"
        "txt_先跌后涨新版5日外=snapshot:20260911"
    )
    assert str(meta["indicator_review_note"]).startswith("snapshot:")
    assert meta["indicator_review_asof"] == ASOF
    # 说明区标注来源=已发布快照信号日（与周五链预计算区分）
    assert "已发布快照信号日" in [
        wb["735"].cell(r, 2).value for r in range(1, 5)
    ] or any(
        "快照" in str(wb["735"].cell(r, 2).value or "")
        for r in range(1, 5)
    )


# ---------------------------------------------------------------------------
# 2. 部分覆盖 → 混合来源，uncovered 走现算
# ---------------------------------------------------------------------------


def test_partial_coverage_mixed_sources(monkeypatch, tmp_path):
    """快照只含 735：covered 用快照，user_demo（不在快照）走即时计算；
    两张 sheet 都在、来源标注正确，导出不因快照缺规则失败。"""
    cfg = _export_env(monkeypatch, tmp_path)
    _publish_snapshot(
        cfg, rules=[("txt_735金叉及趋势", "735", ("SZSE.000001.SZ",))]
    )

    called = {}

    def _fake_compute(cfg_, asof, rule_ids, *, codes=None, on_progress=None):
        called["rule_ids"] = list(rule_ids)
        called["codes"] = list(codes or [])
        return {
            "asof": asof,
            "status": "ok",
            "rules": [
                {
                    "rule_id": "user_demo",
                    "sheet": "我的规则",
                    "count": 1,
                    "matched": [{"code": "SZSE.000002.SZ", "close": 5.9}],
                }
            ],
        }

    monkeypatch.setattr(bq, "_compute_rules_for_export", _fake_compute)
    path = bq.export_bagua_multi_period_xlsx(
        cfg, date=ASOF, periods=["WEEK", "MONTH"],
        adjust="tushare_qfq", all_stocks=True,
        review_rules=["txt_735金叉及趋势", "user_demo"],
    )
    wb = _load_wb(path)
    # 快照部分：来自快照的命中
    assert [r[0] for r in data_rows(wb["735"])] == ["000001"]
    # 现算部分：mock 返回命中 000002（∩ 票池）
    assert [r[0] for r in data_rows(wb["我的规则"])] == ["000002"]
    # 只把快照没覆盖的规则传给即时计算；codes 传导出票池
    assert called["rule_ids"] == ["user_demo"]
    assert called["codes"] == ["SZSE.000001.SZ", "SZSE.000002.SZ"]
    meta = _meta_of(wb)
    sources = str(meta["indicator_review_rule_sources"])
    assert "txt_735金叉及趋势=snapshot:20260911" in sources
    assert "user_demo=computed:20260911" in sources
    assert meta["indicator_review_snapshot_id"]  # 用了快照就记录身份
    note = str(meta["indicator_review_note"])
    assert note.startswith("snapshot:")
    assert "快照未覆盖 1 条" in note
    assert "即时计算" in note


def test_stale_rule_fingerprint_falls_to_compute(monkeypatch, tmp_path):
    """快照指纹与当前公式不符（stale）：该规则不得用快照顶替，走现算。

    构造：发布快照后**改写**快照内的指纹字段模拟公式变更（文件不可变，
    这里直接发布前篡改 payload 指纹，等效于「快照时公式 ≠ 当前公式」）。
    """
    cfg = _export_env(monkeypatch, tmp_path)
    from wtpy.apps.astock.service.screening import _current_rule_fingerprints

    real_fps = _current_rule_fingerprints(cfg, ["txt_735金叉及趋势"])

    class _TamperedFps(dict):
        def __init__(self):
            super().__init__(real_fps or {"txt_735金叉及趋势": "x"})
            self["txt_735金叉及趋势"] = "STALE_FP"

    orig = ss.write_and_publish_snapshot

    def _publish_tampered(cfg_, payload, **kw):
        payload = dict(payload)
        payload["rule_fingerprints"] = _TamperedFps()
        return orig(cfg_, payload, **kw)

    monkeypatch.setattr(ss, "write_and_publish_snapshot", _publish_tampered)
    _publish_snapshot(
        cfg, rules=[("txt_735金叉及趋势", "735", ("SZSE.000001.SZ",))]
    )

    called = {}

    def _fake_compute(cfg_, asof, rule_ids, *, codes=None, on_progress=None):
        called["rule_ids"] = list(rule_ids)
        return {
            "asof": asof, "status": "ok",
            "rules": [{"rule_id": "txt_735金叉及趋势", "sheet": "735",
                       "count": 1,
                       "matched": [{"code": "SZSE.000001.SZ", "close": 1.0}]}],
        }

    monkeypatch.setattr(bq, "_compute_rules_for_export", _fake_compute)
    path = bq.export_bagua_multi_period_xlsx(
        cfg, date=ASOF, adjust="tushare_qfq", all_stocks=True,
        review_rules=["txt_735金叉及趋势"],
    )
    wb = _load_wb(path)
    assert called["rule_ids"] == ["txt_735金叉及趋势"]  # stale → 现算
    meta = _meta_of(wb)
    assert str(meta["indicator_review_rule_sources"]) == (
        "txt_735金叉及趋势=computed:20260911"
    )
    # 快照一个规则都没覆盖上 → 不算「用了快照」，无 snapshot_id
    assert meta["indicator_review_snapshot_id"] in ("", None)


# ---------------------------------------------------------------------------
# 3. 无快照 / 快照不健康 → 行为与旧路径完全一致
# ---------------------------------------------------------------------------


def _export_with_review(monkeypatch, tmp_path, review_rules, *, with_snapshot):
    """同一环境跑两次导出：无快照 vs 有（不健康的）快照，输出必须一致。

    不健康构造：week_id 与 asof 不符（归属校验失败 → 当作无快照）。
    旧路径基线：写一份 review_{asof}.json（带真实指纹）供 load_review_for_export。
    """
    cfg = _export_env(monkeypatch, tmp_path)

    from wtpy.apps.astock.service import indicator_review as ir
    from tests.apps.astock.conftest import isolated_indicator_dir

    # 与 test_bagua_query._review_payload 同口径：真实指纹的 review JSON
    ind_dir = isolated_indicator_dir(tmp_path / "ifp")
    fp_cfg = SimpleNamespace(
        indicator_dir=ind_dir,
        mapping_path=Path(ind_dir) / "_no_tn6_map.json",
        storage_root=tmp_path / "ifp_st",
    )
    try:
        specs = ir._resolve_rules_for_fingerprint(
            fp_cfg, list(ir.DEFAULT_REVIEW_RULES)
        )
        fps = {rid: ir._spec_fingerprint(sp) for rid, _s, sp in specs}
    except Exception:  # noqa: BLE001  CI 公式 fixture 缺失时无法比对 → 跳过
        pytest.skip("公式指纹源不可用，无法构造 review 基线")

    payload = {
        "asof": ASOF,
        "generated_at": "2026-09-11 19:00:00",
        "status": "ok",
        "no_go_reason": "",
        "universe_size": 2,
        "scanned": 2,
        "error_count": 0,
        "rules": [
            {"rule_id": "txt_735金叉及趋势", "sheet": "735", "count": 1,
             "matched": [{"code": "SZSE.000001.SZ", "close": 5.9}]},
        ],
        ir._FP_RULE: fps,
        ir._FP_UNIVERSE: "u",
        ir._FP_NAME: "n",
        ir._FP_SURFACE: "s",
    }
    d = tmp_path / "indicator_review"
    d.mkdir(parents=True, exist_ok=True)
    (d / f"review_{ASOF}.json").write_text(
        __import__("json").dumps(payload, ensure_ascii=False), encoding="utf-8"
    )

    if with_snapshot:
        # 发布一个 week_id 归属不符的快照：_load_export_snapshot_for_week
        # 必须把它当「无快照」处理，导出落回旧路径。
        from wtpy.apps.astock.service.screening import _current_rule_fingerprints

        real_fps2 = _current_rule_fingerprints(cfg, ["txt_735金叉及趋势"])
        bad = {
            "schema_version": ss.SNAPSHOT_SCHEMA_VERSION,
            "snapshot_id": sc.new_snapshot_id(ASOF),
            "run_kind": "weekly_chain",
            "week_id": 20260904,  # 归属不符（asof=20260911）
            "asof": ASOF,
            "generated_at": "2026-09-11 18:40:00",
            "status": "ok",
            "universe_size": len(UNIVERSE),
            "universe_codes": list(UNIVERSE),
            "universe_fingerprint": "ufp1",
            "name_snapshot_id": "ns1",
            "rule_fingerprints": real_fps2 or {"txt_735金叉及趋势": "fp"},
            "data_version": {},
            "content_fingerprint": "cfp1",
            "scanned": len(UNIVERSE),
            "missing_count": 0,
            "no_data_codes": [],
            "rules": [
                {"rule_id": "txt_735金叉及趋势", "sheet": "735", "status": "ok",
                 "count": 2,
                 "matched": [
                     {"code": "SZSE.000001.SZ", "close": 1.0},
                     {"code": "SZSE.000002.SZ", "close": 1.0},
                 ],
                 "failed_codes": []},
            ],
            "duration_sec": 0.1,
        }
        pub = ss.write_and_publish_snapshot(cfg, bad)
        assert pub["published"] is True

    return bq.export_bagua_multi_period_xlsx(
        cfg, date=ASOF, periods=["WEEK", "MONTH"],
        adjust="tushare_qfq", all_stocks=True,
        review_rules=review_rules,
    )


def test_no_snapshot_matches_legacy_path(monkeypatch, tmp_path):
    """无快照：显式勾选走旧路径（review JSON precomputed），meta 无 snapshot_id。"""
    path = _export_with_review(
        monkeypatch, tmp_path, ["txt_735金叉及趋势"], with_snapshot=False
    )
    wb = _load_wb(path)
    assert "735" in wb.sheetnames
    assert [r[0] for r in data_rows(wb["735"])] == ["000001"]
    meta = _meta_of(wb)
    assert meta["indicator_review_snapshot_id"] in ("", None)
    assert str(meta["indicator_review_rule_sources"]) == (
        "txt_735金叉及趋势=precomputed:20260911"
    )
    assert meta["indicator_review_note"] == "ok"


def test_unhealthy_snapshot_falls_back_like_legacy(monkeypatch, tmp_path):
    """快照 week 归属不符：当无快照处理，输出与旧路径逐 sheet 一致。

    归属不符的快照里 735 多命中了 000002（若被误用名单会不同）——用它
    反证「坏快照绝不被读取」。
    """
    path = _export_with_review(
        monkeypatch, tmp_path, ["txt_735金叉及趋势"], with_snapshot=True
    )
    wb = _load_wb(path)
    # 命中必须来自 review JSON（只有 000001），而不是坏快照（000001+000002）
    assert [r[0] for r in data_rows(wb["735"])] == ["000001"]
    meta = _meta_of(wb)
    assert meta["indicator_review_snapshot_id"] in ("", None)
    assert str(meta["indicator_review_rule_sources"]) == (
        "txt_735金叉及趋势=precomputed:20260911"
    )


def test_status_not_ok_snapshot_ignored(monkeypatch, tmp_path):
    """快照 status=no_go（发布门槛外的历史产物）：当无快照。"""
    cfg = _export_env(monkeypatch, tmp_path)
    # 直接手工造一个「指针指向 no_go 快照」的状态（绕过发布门槛）
    snap_dir = tmp_path / "screen_snapshots"
    snap_dir.mkdir(parents=True, exist_ok=True)
    bad_id = sc.new_snapshot_id(ASOF)
    (snap_dir / f"snap_{bad_id}.json").write_text(
        __import__("json").dumps({
            "snapshot_id": bad_id, "run_kind": "weekly_chain",
            "week_id": ASOF, "asof": ASOF, "status": "no_go",
            "universe_codes": [], "rules": [], "rule_fingerprints": {},
        }, ensure_ascii=False),
        encoding="utf-8",
    )
    idx = {"schema_version": "1", "weeks": {
        str(ASOF): {"published_snapshot_id": bad_id, "run_kind": "weekly_chain"}
    }}
    (snap_dir / "index.json").write_text(
        __import__("json").dumps(idx, ensure_ascii=False), encoding="utf-8"
    )

    called = {}

    def _fake_compute(cfg_, asof, rule_ids, *, codes=None, on_progress=None):
        called["rule_ids"] = list(rule_ids)
        return {"asof": asof, "status": "ok", "rules": []}

    monkeypatch.setattr(bq, "_compute_rules_for_export", _fake_compute)
    path = bq.export_bagua_multi_period_xlsx(
        cfg, date=ASOF, adjust="tushare_qfq", all_stocks=True,
        review_rules=["txt_735金叉及趋势"],
    )
    wb = _load_wb(path)
    meta = _meta_of(wb)
    # status 非 ok → 当无快照：无 snapshot_id、全量走现算（无 review 文件）
    assert meta["indicator_review_snapshot_id"] in ("", None)
    assert called["rule_ids"] == ["txt_735金叉及趋势"]


# ---------------------------------------------------------------------------
# 4. 一致性：导出命中∩票池 == combine_snapshot_hits 同票池 hits
# ---------------------------------------------------------------------------


def test_export_hits_equal_combine_snapshot_hits(monkeypatch, tmp_path):
    """验收口径：同快照同规则同票池，导出信号 sheet 命中与筛选接口
    （combine_snapshot_hits）一致——两者读同一快照文件、同做票池交集。"""
    cfg = _export_env(monkeypatch, tmp_path)
    # 快照里 735 命中 3 只中的 000001/000003；导出票池只有 000001/000002
    _publish_snapshot(
        cfg,
        rules=[(
            "txt_735金叉及趋势", "735",
            ("SZSE.000001.SZ", "SZSE.000003.SZ"),
        )],
    )
    path = bq.export_bagua_multi_period_xlsx(
        cfg, date=ASOF, adjust="tushare_qfq", all_stocks=True,
        review_rules=["txt_735金叉及趋势"],
    )
    wb = _load_wb(path)
    export_hits = [r[0] for r in data_rows(wb["735"])]

    # 筛选接口路径：同快照（按 published 指针重读）、同规则、any 模式、
    # 范围=导出票池 → combine_snapshot_hits 的 hits
    snap = ss.load_published_snapshot_for_week(cfg, ASOF)
    assert snap is not None
    pool = ["SZSE.000001.SZ", "SZSE.000002.SZ"]
    combined = ss.combine_snapshot_hits(
        snap, rule_ids=["txt_735金叉及趋势"], match_mode="any", codes=pool
    )
    screen_hits = [h["code"] for h in combined["hits"]]

    # 导出展示 6 位码、筛选接口标准码：按 6 位码对齐后做集合级一致断言
    # （标准码 SZSE.000001.SZ 的 6 位码是中段）
    export_six = set(export_hits)
    screen_six = {c.split(".")[1] for c in screen_hits}
    assert export_six == screen_six == {"000001"}


# ---------------------------------------------------------------------------
# 5. review_rules=None 兼容层：不进快照路径
# ---------------------------------------------------------------------------


def test_review_rules_none_never_reads_snapshot(monkeypatch, tmp_path):
    """review_rules=None（CLI/旧调用）：即便有可用快照也走 review JSON 全量
    ——快照路径不得改变兼容层的输出规则集（评审第 7 条双路径分离）。"""
    cfg = _export_env(monkeypatch, tmp_path)
    _publish_snapshot(cfg)  # 快照含两条规则且可覆盖

    # 旧路径基线：写一条只含 735 的 review JSON（与快照规则集不同，
    # 用于证明 None 路径读的是 JSON 而非快照）
    from wtpy.apps.astock.service import indicator_review as ir

    ind_dir = cfg.indicator_dir
    fp_cfg = SimpleNamespace(
        indicator_dir=ind_dir, mapping_path=cfg.mapping_path,
        storage_root=tmp_path,
    )
    try:
        specs = ir._resolve_rules_for_fingerprint(
            fp_cfg, [("txt_735金叉及趋势", "735")]
        )
        fps = {rid: ir._spec_fingerprint(sp) for rid, _s, sp in specs}
    except Exception:  # noqa: BLE001
        pytest.skip("公式指纹源不可用")

    import json as _json

    d = tmp_path / "indicator_review"
    d.mkdir(parents=True, exist_ok=True)
    (d / f"review_{ASOF}.json").write_text(
        _json.dumps({
            "asof": ASOF, "generated_at": "2026-09-11 19:00:00",
            "status": "ok", "no_go_reason": "", "universe_size": 2,
            "scanned": 2, "error_count": 0,
            "rules": [
                {"rule_id": "txt_735金叉及趋势", "sheet": "735", "count": 1,
                 "matched": [{"code": "SZSE.000001.SZ", "close": 5.9}]},
            ],
            ir._FP_RULE: fps, ir._FP_UNIVERSE: "u",
            ir._FP_NAME: "n", ir._FP_SURFACE: "s",
        }, ensure_ascii=False),
        encoding="utf-8",
    )

    path = bq.export_bagua_multi_period_xlsx(
        cfg, date=ASOF, adjust="tushare_qfq", all_stocks=True,
        review_rules=None,
    )
    wb = _load_wb(path)
    # None 路径输出 = review JSON 的规则集（只有 735），不是快照的两条
    assert "735" in wb.sheetnames
    assert "5日外" not in wb.sheetnames
    meta = _meta_of(wb)
    assert meta["indicator_review_snapshot_id"] in ("", None)
    assert str(meta["indicator_review_rule_sources"]) == (
        "txt_735金叉及趋势=precomputed:20260911"
    )


# ---------------------------------------------------------------------------
# 6. 纯函数单测（适配层语义锁定）
# ---------------------------------------------------------------------------


def test_snapshot_rules_to_review_form_filters_error_rules():
    """适配层：status=error 规则被剔除（缺规则 ≠ 零命中，不得端出），
    其余规则字段原样转换（matched 的 code/close 保真）。"""
    snap = {
        "asof": 20260911, "status": "ok",
        "rules": [
            {
                "rule_id": "r_ok", "sheet": "甲", "status": "ok", "count": 1,
                "matched": [{"code": "SZSE.000001.SZ", "close": 10.5}],
                "failed_codes": [],
            },
            {
                "rule_id": "r_partial", "sheet": "乙", "status": "partial",
                "count": 1,
                "matched": [{"code": "SZSE.000002.SZ", "close": None}],
                "failed_codes": ["SZSE.000003.SZ"],
            },
            {
                "rule_id": "r_err", "sheet": "丙", "status": "error",
                "count": 0, "matched": [], "failed_codes": [],
            },
        ],
    }
    out = bq._snapshot_rules_to_review_form(snap)
    assert out["asof"] == 20260911
    ids = [r["rule_id"] for r in out["rules"]]
    assert "r_err" not in ids  # error 规则不进入导出形态
    assert set(ids) == {"r_ok", "r_partial"}  # partial 保留（hit/miss 可用）
    r_ok = next(r for r in out["rules"] if r["rule_id"] == "r_ok")
    assert r_ok["matched"] == [{"code": "SZSE.000001.SZ", "close": 10.5}]


def test_load_export_snapshot_for_week_health_checks(monkeypatch, tmp_path):
    """快照读取预检：status 非 ok / week_id / asof 不符 → 一律 None
    （当无快照，不算错误）；正常快照原样返回。"""
    cfg = SimpleNamespace(storage_root=tmp_path)
    snap_dir = tmp_path / "screen_snapshots"
    snap_dir.mkdir(parents=True, exist_ok=True)

    def _put(snapshot_id, payload, week_id=None):
        (snap_dir / f"snap_{snapshot_id}.json").write_text(
            __import__("json").dumps(payload, ensure_ascii=False),
            encoding="utf-8",
        )
        idx = {"schema_version": "1", "weeks": {
            str(week_id or payload.get("week_id")): {
                "published_snapshot_id": snapshot_id,
                "run_kind": "weekly_chain",
            }
        }}
        (snap_dir / "index.json").write_text(
            __import__("json").dumps(idx, ensure_ascii=False), encoding="utf-8"
        )

    good_id = "good_001"
    _put(good_id, {
        "snapshot_id": good_id, "week_id": ASOF, "asof": ASOF,
        "status": "ok", "rules": [], "rule_fingerprints": {},
    })
    snap = bq._load_export_snapshot_for_week(cfg, ASOF)
    assert snap is not None and snap["snapshot_id"] == good_id

    bad_status = "bad_002"
    _put(bad_status, {
        "snapshot_id": bad_status, "week_id": ASOF, "asof": ASOF,
        "status": "no_go", "rules": [], "rule_fingerprints": {},
    })
    assert bq._load_export_snapshot_for_week(cfg, ASOF) is None

    bad_week = "bad_003"
    _put(bad_week, {
        "snapshot_id": bad_week, "week_id": 20260904, "asof": ASOF,
        "status": "ok", "rules": [], "rule_fingerprints": {},
    })
    assert bq._load_export_snapshot_for_week(cfg, ASOF) is None

    # 无发布指针 → None（load_published_snapshot_for_week 契约）
    (snap_dir / "index.json").write_text(
        '{"schema_version": "1", "weeks": {}}', encoding="utf-8"
    )
    assert bq._load_export_snapshot_for_week(cfg, ASOF) is None


class TestPartialRuleVisibility:
    """审查 C-1 回归锁：partial 规则的非全量性必须在导出侧可见。"""

    def test_snapshot_partial_rule_carries_warning(self):
        """适配层保留 partial 状态与失败票数（sheet 说明区/meta 可据此标注）。"""
        from wtpy.apps.astock.service.bagua_query import _snapshot_rules_to_review_form

        snap = {
            "asof": 20260911,
            "generated_at": "2026-09-11 18:40:00",
            "status": "ok",
            "rules": [
                {"rule_id": "r_ok", "sheet": "完整规则", "status": "ok",
                 "count": 1, "matched": [{"code": "SZSE.STK.000001", "close": 10.0}],
                 "failed_codes": []},
                {"rule_id": "r_partial", "sheet": "部分规则", "status": "partial",
                 "count": 1, "matched": [{"code": "SZSE.STK.000002", "close": 20.0}],
                 "failed_codes": [f"SZSE.STK.0000{i}" for i in range(100, 104)]},
                {"rule_id": "r_error", "sheet": "失败规则", "status": "error",
                 "count": 0, "matched": [], "failed_codes": ["SZSE.STK.000009"]},
            ],
        }
        out = _snapshot_rules_to_review_form(snap)
        by_id = {r["rule_id"]: r for r in out["rules"]}
        # error 剔除（不冒充零命中）
        assert "r_error" not in by_id
        # partial 保留 + 状态与失败数透传（下游据此打完整性提示）
        assert by_id["r_partial"]["snapshot_status"] == "partial"
        assert by_id["r_partial"]["snapshot_failed_count"] == 4
        # ok 规则不带 partial 标注
        assert by_id["r_ok"]["snapshot_status"] == "ok"
        assert by_id["r_ok"]["snapshot_failed_count"] == 0

    def test_rule_brief_rows_renders_partial_note(self, tmp_path):
        """sheet 说明区渲染「完整性提示」行（_rule_brief_rows 接 note）。"""
        from wtpy.apps.astock.service.bagua_query import _rule_brief_rows
        from wtpy.apps.astock.config import get_default_config

        cfg = get_default_config(storage_root=tmp_path / "st")
        brief = {
            "rule_id": "txt_x",
            "source_kind": "snapshot",
            "asof": 20260911,
            "count": 5,
            "raw_count": 5,
            "note": "注意：该规则本次评估有 4 只股票未完成（失败/缺数据），命中名单非全量",
        }
        rows = _rule_brief_rows(cfg, brief)
        notes = [v for k, v in rows if k == "完整性提示"]
        assert notes and "非全量" in notes[0]
