# -*- coding: utf-8 -*-
"""独立端到端验证（三项功能回归）：
A. 回测并发可配置（HTTP + JobStore resize + 持久化 + 老库缺表容忍）
B. calendar_range 的 data_max_date（真实数据最后可用日）
C. 导出信号 sheet 的复核基准日回退（service / 同步响应头 / 异步 job 字段）
D. 前端静态契约（#btQueueWorkers、data_max_date、bqHint 回退文案）

与 coder 已加用例（test_job_queue.py / test_ui_universe_dates.py /
test_bagua_query.py）独立编写，仅验证跨层集成表现，不改业务代码。
"""
from __future__ import annotations

import datetime as _dt
import io
import json
import re
import threading
import time
from pathlib import Path
from unittest.mock import patch
from urllib.parse import unquote

import pytest

import tests.apps.astock.conftest  # noqa: F401

from wtpy.apps.astock.config import AStockConfig, get_default_config
from wtpy.apps.astock.service.backtest import BacktestRequest
from wtpy.apps.astock.service.db import (
    connect,
    get_app_setting,
    set_app_setting,
)
from wtpy.apps.astock.service.jobs import (
    DEFAULT_BT_MAX_WORKERS,
    HARD_MAX_BT_WORKERS,
    JobStore,
)

BACKEND_ROOT = Path(__file__).resolve().parents[3]
STATIC_DIR = BACKEND_ROOT / "wtpy" / "apps" / "astock" / "web" / "static"
V3_HTML = STATIC_DIR / "index_v3.html"
BAGUA_JSON = (
    BACKEND_ROOT / "wtpy" / "apps" / "astock" / "bagua" / "bagua_384.json"
)

_DS_META_MOCK = {
    "dataset_id": "mock",
    "dataset_source": "tdxquant",
    "dataset_adjustment": "front",
    "dataset_status": "ready",
    "covers_asof": True,
    "candidate_datasets": 1,
}


# ---------------------------------------------------------------------------
# 公共工具
# ---------------------------------------------------------------------------


def _make_app(cfg: AStockConfig):
    from wtpy.apps.astock.api import create_app

    return create_app(cfg)


def _make_client(app):
    from fastapi.testclient import TestClient

    return TestClient(app)


def _wait_alive(store: JobStore, n: int, timeout: float = 3.0):
    deadline = time.time() + timeout
    alive = [t for t in store._workers if t.is_alive()]
    while time.time() < deadline:
        alive = [t for t in store._workers if t.is_alive()]
        if len(alive) == n:
            return alive
        time.sleep(0.02)
    return alive


def _wait_status(store: JobStore, job_id: str, statuses, timeout: float = 5.0) -> str:
    if isinstance(statuses, str):
        statuses = {statuses}
    else:
        statuses = set(statuses)
    deadline = time.time() + timeout
    last = ""
    while time.time() < deadline:
        last = store.get(job_id).status
        if last in statuses:
            return last
        time.sleep(0.02)
    return last


def _weekdays(start: str, end: str) -> list:
    out = []
    cur = _dt.date.fromisoformat(start)
    end_d = _dt.date.fromisoformat(end)
    while cur <= end_d:
        if cur.isoweekday() <= 5:
            out.append(int(cur.strftime("%Y%m%d")))
        cur += _dt.timedelta(days=1)
    return out


def _extract_js_function(src: str, name: str) -> str:
    m = re.search(r"(?:async\s+)?function\s+" + name + r"\s*\(", src)
    assert m, f"function not found: {name}"
    brace = src.index("{", m.start())
    depth = 0
    for j in range(brace, len(src)):
        if src[j] == "{":
            depth += 1
        elif src[j] == "}":
            depth -= 1
            if depth == 0:
                return src[m.start(): j + 1]
    raise AssertionError(f"unbalanced {name}")


# ---------------------------------------------------------------------------
# A. 回测并发可配置
# ---------------------------------------------------------------------------


@pytest.fixture()
def api_cfg(tmp_path: Path) -> AStockConfig:
    storage = tmp_path / "st"
    ind = tmp_path / "ind"
    storage.mkdir()
    ind.mkdir()
    return get_default_config(
        storage_root=storage,
        indicator_dir=ind,
        output_root=tmp_path / "out",
    )


def test_queue_config_default_put_persist_restart_and_400(
    api_cfg: AStockConfig, monkeypatch
):
    """GET 默认结构 → PUT 3（db source / 生效 / 落盘）→ 新 app 重启仍 3；
    非法值 400 且不清空既有配置。"""
    pytest.importorskip("fastapi")
    monkeypatch.delenv("ASTOCK_BT_MAX_WORKERS", raising=False)

    app = _make_app(api_cfg)
    client = _make_client(app)
    jobs = app.state.astock.jobs
    try:
        d = client.get("/api/v1/backtests/queue/config").json()
        assert set(d) >= {
            "max_workers",
            "hard_max_workers",
            "default_workers",
            "env_override",
            "source",
        }
        assert d["max_workers"] == DEFAULT_BT_MAX_WORKERS == 6
        assert d["default_workers"] == 6
        assert d["hard_max_workers"] == HARD_MAX_BT_WORKERS == 8
        assert d["env_override"] is None
        assert d["source"] == "default"

        r = client.put("/api/v1/backtests/queue/config", json={"max_workers": 3})
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["max_workers"] == 3
        assert d["source"] == "db"
        assert jobs.max_workers == 3
        assert len(_wait_alive(jobs, 3)) == 3
        assert get_app_setting(api_cfg, "bt_max_workers") == "3"

        for bad in (0, 9, "abc", True):
            rb = client.put(
                "/api/v1/backtests/queue/config", json={"max_workers": bad}
            )
            assert rb.status_code == 400, (bad, rb.status_code, rb.text)
        # 400 之后内存/持久化都不得被破坏
        assert jobs.max_workers == 3
        assert get_app_setting(api_cfg, "bt_max_workers") == "3"
    finally:
        jobs.shutdown(wait=True)

    # 同一 cfg 重启新 app：读持久化，不落回 env/默认
    app2 = _make_app(api_cfg)
    try:
        assert app2.state.astock.jobs.max_workers == 3
        d2 = _make_client(app2).get("/api/v1/backtests/queue/config").json()
        assert d2["max_workers"] == 3
        assert d2["source"] == "db"
    finally:
        app2.state.astock.jobs.shutdown(wait=True)


def test_experiments_presets_concurrency_reflects_config(
    api_cfg: AStockConfig, monkeypatch
):
    """presets.default_concurrency 跟随持久化值；hard_max_concurrency 固定 8。"""
    pytest.importorskip("fastapi")
    monkeypatch.delenv("ASTOCK_BT_MAX_WORKERS", raising=False)
    app = _make_app(api_cfg)
    client = _make_client(app)
    try:
        body = client.get("/api/v1/experiments/presets").json()
        assert body["default_concurrency"] == 6
        assert body["hard_max_concurrency"] == HARD_MAX_BT_WORKERS == 8

        client.put("/api/v1/backtests/queue/config", json={"max_workers": 4})
        body = client.get("/api/v1/experiments/presets").json()
        assert body["default_concurrency"] == 4
        assert body["hard_max_concurrency"] == 8
    finally:
        app.state.astock.jobs.shutdown(wait=True)


def test_shrink_8_to_2_keeps_queue_consuming_and_clean_shutdown(
    api_cfg: AStockConfig, monkeypatch
):
    """PUT 8 扩到 8 → PUT 2 缩到 2：活着线程 == 2，且哨兵不级联（队列仍消费）。"""
    pytest.importorskip("fastapi")
    monkeypatch.delenv("ASTOCK_BT_MAX_WORKERS", raising=False)
    app = _make_app(api_cfg)
    client = _make_client(app)
    jobs = app.state.astock.jobs
    try:
        assert client.put(
            "/api/v1/backtests/queue/config", json={"max_workers": 8}
        ).json()["max_workers"] == 8
        assert len(_wait_alive(jobs, 8)) == 8

        r = client.put("/api/v1/backtests/queue/config", json={"max_workers": 2})
        assert r.status_code == 200 and r.json()["max_workers"] == 2
        alive = _wait_alive(jobs, 2)
        assert len(alive) == 2, "resize-down 必须只退 6 个且不级联清空"
        assert len(jobs._workers) >= 2

        # 仍然能消费队列（若哨兵误清空全部 worker，这里会超时）
        with patch(
            "wtpy.apps.astock.service.jobs.BacktestService.run",
            return_value={"run_id": "bt_v", "status": "ok", "metrics": {}},
        ):
            rec = jobs.submit(BacktestRequest(rule_ids=["verify"], period="DAY"))
            assert _wait_status(jobs, rec.job_id, "succeeded") == "succeeded"
        assert jobs.queue_snapshot()["max_workers"] == 2

        t0 = time.time()
        jobs.shutdown(wait=True)
        assert time.time() - t0 < 10, "缩容后 shutdown 不应挂起"
    finally:
        jobs.shutdown(wait=False)
    assert not [t for t in jobs._workers if t.is_alive()]


def test_put_resize_with_running_job_does_not_interrupt(
    api_cfg: AStockConfig, monkeypatch
):
    """PUT 缩容时运行中的任务不被打断，且最终成功。"""
    pytest.importorskip("fastapi")
    monkeypatch.delenv("ASTOCK_BT_MAX_WORKERS", raising=False)
    app = _make_app(api_cfg)
    client = _make_client(app)
    jobs = app.state.astock.jobs
    started = threading.Event()
    release = threading.Event()

    def fake_run(self, req, progress_cb=None):
        started.set()
        release.wait(5.0)
        return {"run_id": "bt_hold", "status": "ok", "metrics": {}}

    try:
        with patch("wtpy.apps.astock.service.jobs.BacktestService.run", fake_run):
            rec = jobs.submit(BacktestRequest(rule_ids=["hold"], period="DAY"))
            assert started.wait(3.0), "任务未启动"
            r = client.put(
                "/api/v1/backtests/queue/config", json={"max_workers": 2}
            )
            assert r.status_code == 200 and r.json()["max_workers"] == 2
            assert jobs.get(rec.job_id).status == "running", "缩容不得中断运行中任务"
            release.set()
            assert (
                _wait_status(jobs, rec.job_id, "succeeded", timeout=5.0)
                == "succeeded"
            )
    finally:
        release.set()
        jobs.shutdown(wait=False)


def test_set_max_workers_same_value_noop_and_repeat_shrink_to_one(
    tmp_path: Path,
):
    """同值 resize 不增删线程；连续缩到 1 不累积哨兵、仍可消费任务。"""
    cfg = AStockConfig()
    cfg.storage_root = tmp_path / "store"
    cfg.output_root = tmp_path / "out"
    cfg.storage_root.mkdir(parents=True)
    cfg.output_root.mkdir(parents=True)
    store = JobStore(cfg, max_workers=3)
    try:
        before = list(store._workers)
        assert store.set_max_workers(3) == 3
        assert list(store._workers) == before, "同值 resize 不应增删线程"

        assert store.set_max_workers(8) == 8
        assert len(_wait_alive(store, 8)) == 8
        assert store.set_max_workers(1) == 1
        assert len(_wait_alive(store, 1)) == 1
        # 重复同值缩容：哨兵不得遗留/级联（否则 worker 变 0）
        assert store.set_max_workers(1) == 1
        assert len(_wait_alive(store, 1)) == 1

        with patch(
            "wtpy.apps.astock.service.jobs.BacktestService.run",
            return_value={"run_id": "bt_one", "status": "ok", "metrics": {}},
        ):
            rec = store.submit(BacktestRequest(rule_ids=["one"], period="DAY"))
            assert _wait_status(store, rec.job_id, "succeeded") == "succeeded"
    finally:
        store.shutdown(wait=True)
    assert not [t for t in store._workers if t.is_alive()]


def test_missing_app_settings_table_tolerated_and_startup_default(
    api_cfg: AStockConfig, monkeypatch
):
    """老库删掉 app_settings：get_app_setting 不抛、create_app 回默认 6；
    set_app_setting 幂等重建表。"""
    pytest.importorskip("fastapi")
    monkeypatch.delenv("ASTOCK_BT_MAX_WORKERS", raising=False)

    set_app_setting(api_cfg, "bt_max_workers", "3")
    conn = connect(api_cfg)
    conn.execute("DROP TABLE app_settings")
    conn.commit()
    conn.close()

    assert get_app_setting(api_cfg, "bt_max_workers") is None  # 不抛异常

    app = _make_app(api_cfg)
    try:
        assert app.state.astock.jobs.max_workers == 6
        d = _make_client(app).get("/api/v1/backtests/queue/config").json()
        assert d["max_workers"] == 6
        assert d["source"] == "default"
    finally:
        app.state.astock.jobs.shutdown(wait=True)

    set_app_setting(api_cfg, "bt_max_workers", "4")
    assert get_app_setting(api_cfg, "bt_max_workers") == "4"


def test_clamp_concurrency_independent_edges():
    """实验并发 clamp 的独立边界（bool/负值/超上限/非法浮点）。"""
    from wtpy.apps.astock.service.experiments import clamp_concurrency

    assert clamp_concurrency(None) == 1
    assert clamp_concurrency(0) == 1
    assert clamp_concurrency(-5) == 1
    assert clamp_concurrency(8) == 8
    assert clamp_concurrency(99) == HARD_MAX_BT_WORKERS
    assert clamp_concurrency(2.5) == 1
    assert clamp_concurrency(True) == 1
    assert clamp_concurrency("abc") == 1


# ---------------------------------------------------------------------------
# B. calendar_range：data_max_date
# ---------------------------------------------------------------------------


def _publish_ready_dataset(md: Path, *, cutoff: int, last_date: int) -> None:
    import numpy as np

    from wtpy.apps.astock.data.dataset_store import (
        DatasetManifest,
        DatasetStore,
        SymbolRecord,
    )

    store = DatasetStore(md)
    dates = np.array([cutoff], dtype=np.int64) if last_date == cutoff else np.array(
        [cutoff, last_date], dtype=np.int64
    )
    close = np.linspace(10.0, 11.0, len(dates))
    sha = store.store_bar_arrays(
        "SSE.STK.600000",
        {
            "trade_date": dates,
            "open": close,
            "high": close,
            "low": close,
            "close": close,
            "volume": np.ones(len(dates)),
            "amount": np.ones(len(dates)),
        },
    )
    store.publish(
        DatasetManifest(
            dataset_id=f"tushare_none_1d_{cutoff}_verify",
            source="tushare",
            adjustment="none",
            period="1d",
            status="ready",
            data_cutoff_date=cutoff,
            symbols=[
                SymbolRecord(
                    symbol="SSE.STK.600000",
                    blob_sha256=sha,
                    first_date=int(dates[0]),
                    last_date=int(dates[-1]),
                    row_count=len(dates),
                    quality="ok",
                )
            ],
            symbol_count=1,
            row_count=len(dates),
        )
    )


def test_calendar_range_data_max_date_exists_int_and_bounded(
    tmp_path: Path, monkeypatch
):
    pytest.importorskip("fastapi")
    storage = tmp_path / "st"
    storage.mkdir()
    (storage / "calendar.json").write_text(
        json.dumps({"dates": [20240102, 20240630, 20241231]}),
        encoding="utf-8",
    )
    cfg = get_default_config(storage_root=storage)
    client = _make_client(_make_app(cfg))
    body = client.get("/api/v1/calendar/range").json()
    assert isinstance(body["data_max_date"], int)
    assert isinstance(body["max_date"], int)
    assert body["min_date"] <= body["data_max_date"] <= body["max_date"]


def test_calendar_range_data_max_date_takes_cutoff_newer_than_calendar(
    tmp_path: Path, monkeypatch
):
    """数据 cutoff（20241231）晚于日历末日（20240630）：data_max_date 取 cutoff。"""
    pytest.importorskip("fastapi")
    storage = tmp_path / "st"
    storage.mkdir()
    md = tmp_path / "md"
    md.mkdir()
    monkeypatch.setenv("MARKET_DATA_ROOT", str(md))
    (storage / "calendar.json").write_text(
        json.dumps({"dates": [20240102, 20240630]}),
        encoding="utf-8",
    )
    cfg = get_default_config(storage_root=storage)
    _publish_ready_dataset(md, cutoff=20241231, last_date=20241231)

    body = _make_client(_make_app(cfg)).get("/api/v1/calendar/range").json()
    assert isinstance(body["data_max_date"], int)
    assert body["data_max_date"] == 20241231
    # 路由同时把 max_date 扩到最新 cutoff；data_max_date 必须跟数据走
    assert body["max_date"] == 20241231
    assert 2024 in body["years"]


def test_calendar_range_data_max_date_lags_future_calendar(
    tmp_path: Path, monkeypatch
):
    """日历含未来末日（20261231）、数据止于 20260828：两字段语义分离。"""
    pytest.importorskip("fastapi")
    storage = tmp_path / "st"
    storage.mkdir()
    md = tmp_path / "md"
    md.mkdir()
    monkeypatch.setenv("MARKET_DATA_ROOT", str(md))
    (storage / "calendar.json").write_text(
        json.dumps({"dates": [20240102, 20260828, 20261231]}),
        encoding="utf-8",
    )
    cfg = get_default_config(storage_root=storage)
    _publish_ready_dataset(md, cutoff=20260828, last_date=20260828)

    body = _make_client(_make_app(cfg)).get("/api/v1/calendar/range").json()
    assert body["max_date"] == 20261231
    assert body["data_max_date"] == 20260828


# ---------------------------------------------------------------------------
# C. 导出信号 sheet 复核基准日回退
# ---------------------------------------------------------------------------


def _review_payload(asof: int, matched_735=None, matched_5w=None) -> dict:
    def _matched(codes):
        return [{"code": c, "close": 5.9} for c in (codes or [])]

    return {
        "asof": int(asof),
        "generated_at": "2026-08-14 19:00:00",
        "status": "ok",
        "no_go_reason": "",
        "universe_size": 2,
        "scanned": 2,
        "error_count": 0,
        "rules": [
            {
                "rule_id": "txt_735金叉及趋势",
                "sheet": "735",
                "count": len(matched_735 or []),
                "matched": _matched(matched_735),
            },
            {
                "rule_id": "txt_先跌后涨新版5日外",
                "sheet": "5日外",
                "count": len(matched_5w or []),
                "matched": _matched(matched_5w),
            },
        ],
    }


def _write_review(cfg: AStockConfig, asof: int, payload: dict) -> None:
    d = Path(cfg.storage_root) / "indicator_review"
    d.mkdir(parents=True, exist_ok=True)
    (d / f"review_{asof}.json").write_text(
        json.dumps(payload, ensure_ascii=False), encoding="utf-8"
    )


def _mock_export_surface(
    monkeypatch,
    cfg: AStockConfig,
    *,
    surface_max: int,
    review_max_age: int = 30,
):
    """复用 test_bagua_query.py 的 mock 模式：
    - 行情日线走 _load_dataset_bars（数据止于 surface_max）
    - BaguaPlaneSession 不可用 → 数据面桩
    - 正式产品面 max_date = surface_max
    - 放宽复核文件年龄窗口（默认 7 天），隔离验证回退链本身
    """
    from wtpy.apps.astock.data.tdx_reader import DayBar
    from wtpy.apps.astock.service import bagua_query as bq
    from wtpy.apps.astock.service import indicator_review as ir

    days = _weekdays("2026-06-01", "2026-08-28")
    days = [d for d in days if d <= surface_max]

    def _fake_load(cfg_, std_code, source_key, asof=None, **_kw):
        sel = [d for d in days if asof is None or d <= int(asof)]
        return (
            [DayBar(d, 6.27, 7.33, 5.90, 5.90, 1.0, 1.0) for d in sel],
            dict(_DS_META_MOCK),
        )

    monkeypatch.setattr(bq, "_load_dataset_bars", _fake_load)
    monkeypatch.setattr(
        bq,
        "BaguaPlaneSession",
        lambda *_a, **_k: (_ for _ in ()).throw(FileNotFoundError("no md")),
    )
    monkeypatch.setattr(
        bq,
        "_resolve_batch_codes",
        lambda cfg_, codes=None, *, all_stocks=False: [
            "SSE.STK.600000",
            "SSE.STK.000001",
        ],
    )
    monkeypatch.setattr(bq, "list_etf_std_codes", lambda cfg_: [])
    monkeypatch.setattr(
        ir,
        "_resolve_formal_surface",
        lambda _cfg: ({"formal_l1_id": "mock_l1", "max_date": surface_max}, ""),
    )
    _orig_load = ir.load_review_for_export
    monkeypatch.setattr(
        ir,
        "load_review_for_export",
        lambda cfg_, asof, **kw: _orig_load(
            cfg_, asof, max_age_days=review_max_age
        ),
    )
    return cfg


def _export_cfg(tmp_path: Path) -> AStockConfig:
    cfg = get_default_config(
        storage_root=tmp_path / "st",
        indicator_dir=tmp_path / "ind",
        output_root=tmp_path / "out",
    )
    Path(cfg.storage_root).mkdir(parents=True, exist_ok=True)
    Path(cfg.indicator_dir).mkdir(parents=True, exist_ok=True)
    return cfg


def test_export_signal_sheet_uses_review_file_after_surface_fallback(
    tmp_path: Path, monkeypatch
):
    """请求 20260910 超出数据面 20260828，复核文件 20260814：
    信号 sheet 成员来自复核文件、meta/info_out/列头基准日一致。"""
    if not BAGUA_JSON.exists():
        pytest.skip("bagua_384.json missing")
    cfg = _export_cfg(tmp_path)
    _mock_export_surface(monkeypatch, cfg, surface_max=20260828)
    _write_review(cfg, 20260814, _review_payload(20260814, ["SSE.STK.600000"], []))

    from wtpy.apps.astock.service import bagua_query as bq

    info: dict = {}
    path = bq.export_bagua_multi_period_xlsx(
        cfg,
        date="2026-09-10",
        periods=["WEEK", "MONTH"],
        adjust="tushare_qfq",
        all_stocks=True,
        info_out=info,
    )

    import openpyxl

    wb = openpyxl.load_workbook(path)
    assert "735" in wb.sheetnames, "复核文件在窗口内时必须出信号 sheet"
    rows = list(wb["735"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == ["600000"], "成员必须来自复核文件命中"
    assert rows[0][2] == "2026-08-14"

    # 成员行/行内卦象/周列头都取复核文件日（周标签一致）
    week_label = bq._week_iso_label(20260814)
    headers_735 = [c.value for c in wb["735"][1]]
    assert headers_735[8] == f"周卦周线-组合({week_label})"
    # 0 命中的空信号 sheet 列头也必须用回退日，而非请求日所在周
    headers_5w = [c.value for c in wb["5日外"][1]]
    assert headers_5w[8] == f"周卦周线-组合({week_label})"

    meta = {r[0]: r[1] for r in wb["meta"].iter_rows(min_row=2, values_only=True)}
    assert meta["indicator_review_query_date"] == 20260910
    assert meta["indicator_review_asof"] == 20260814
    note = str(meta["indicator_review_note"])
    assert "fallback_date:请求 20260910 超出数据覆盖 20260828" in note
    assert "20260814" in note
    assert meta["indicator_review_sheets"] == "735,5日外"

    assert info["query_date"] == 20260910
    assert info["review_asof_used"] == 20260814
    assert info["review_fallback"] is True
    assert info["review_note"] == note


def test_export_signal_sheet_stale_policy_observation(
    tmp_path: Path, monkeypatch
):
    """记录默认 7 天窗口的策略行为：20260814 复核距数据面 20260828 已 14 天，
    被判 stale → 不生成信号 sheet（回退链本身见上一个用例）。"""
    if not BAGUA_JSON.exists():
        pytest.skip("bagua_384.json missing")
    cfg = _export_cfg(tmp_path)
    _mock_export_surface(monkeypatch, cfg, surface_max=20260828, review_max_age=7)
    _write_review(cfg, 20260814, _review_payload(20260814, ["SSE.STK.600000"], []))

    from wtpy.apps.astock.service import bagua_query as bq

    info: dict = {}
    path = bq.export_bagua_multi_period_xlsx(
        cfg,
        date="2026-09-10",
        periods=["WEEK", "MONTH"],
        adjust="tushare_qfq",
        all_stocks=True,
        info_out=info,
    )

    import openpyxl

    wb = openpyxl.load_workbook(path)
    assert "735" not in wb.sheetnames and "5日外" not in wb.sheetnames
    meta = {r[0]: r[1] for r in wb["meta"].iter_rows(min_row=2, values_only=True)}
    assert str(meta["indicator_review_note"]).startswith("stale")
    # 即便没有信号 sheet，回退基准日仍写 meta（供 UI/上游排查）
    assert meta["indicator_review_query_date"] == 20260910
    assert info["query_date"] == 20260910


def test_bagua_export_sync_route_review_headers(tmp_path: Path, monkeypatch):
    """同步导出路由：X-Bagua-Review-* 响应头存在且 note 可解码。"""
    pytest.importorskip("fastapi")
    if not BAGUA_JSON.exists():
        pytest.skip("bagua_384.json missing")
    cfg = _export_cfg(tmp_path)
    _mock_export_surface(monkeypatch, cfg, surface_max=20260828)
    _write_review(cfg, 20260814, _review_payload(20260814, ["SSE.STK.600000"], []))

    app = _make_app(cfg)
    try:
        client = _make_client(app)
        r = client.post(
            "/api/v1/bagua/export?async_mode=false",
            json={
                "codes": ["600000"],
                "all_stocks": False,
                "date": "2026-09-10",
                "periods": ["WEEK", "MONTH"],
                "adjust": "tushare_qfq",
            },
        )
        assert r.status_code == 200, r.text
        assert r.headers.get("X-Bagua-Review-AsOf") == "20260814"
        assert r.headers.get("X-Bagua-Review-Fallback") == "1"
        note = unquote(r.headers.get("X-Bagua-Review-Note") or "")
        assert "fallback_date:请求 20260910 超出数据覆盖 20260828" in note
        # 响应体仍是可打开的 xlsx，且信号 sheet 非空
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        assert "735" in wb.sheetnames
        rows = list(wb["735"].iter_rows(min_row=2, values_only=True))
        assert [row[0] for row in rows] == ["600000"]
    finally:
        app.state.astock.jobs.shutdown(wait=False)


def test_bagua_export_async_job_review_fields(tmp_path: Path, monkeypatch):
    """异步 job：完成后 query_date / review_asof_used / review_fallback /
    review_note 字段齐全且值正确。"""
    pytest.importorskip("fastapi")
    if not BAGUA_JSON.exists():
        pytest.skip("bagua_384.json missing")
    cfg = _export_cfg(tmp_path)
    _mock_export_surface(monkeypatch, cfg, surface_max=20260828)
    _write_review(cfg, 20260814, _review_payload(20260814, ["SSE.STK.600000"], []))

    app = _make_app(cfg)
    try:
        client = _make_client(app)
        r = client.get(
            "/api/v1/bagua/export",
            params={
                "date": "2026-09-10",
                "period": "WEEK,MONTH",
                "adjust": "tushare_qfq",
                "all_stocks": "true",
                "async_mode": "true",
            },
        )
        assert r.status_code == 200, r.text
        initial = r.json()
        assert initial.get("job_id")
        assert initial.get("mode") == "async"
        for key in (
            "job_id",
            "status",
            "date",
            "periods",
            "adjust",
            "all_stocks",
            "query_date",
            "review_asof_used",
            "review_fallback",
            "review_note",
        ):
            assert key in initial, key

        deadline = time.time() + 30.0
        st = initial
        while time.time() < deadline:
            st = client.get(
                f"/api/v1/bagua/export/jobs/{initial['job_id']}"
            ).json()
            if st.get("status") in ("done", "error"):
                break
            time.sleep(0.1)
        assert st["status"] == "done", st
        assert st["query_date"] == 20260910
        assert st["review_asof_used"] == 20260814
        assert st["review_fallback"] is True
        assert "fallback_date" in st["review_note"]
        assert st.get("filename")
    finally:
        app.state.astock.jobs.shutdown(wait=False)


# ---------------------------------------------------------------------------
# D. 前端静态契约
# ---------------------------------------------------------------------------


def _v3_html() -> str:
    assert V3_HTML.is_file(), f"missing {V3_HTML}"
    return V3_HTML.read_text(encoding="utf-8")


def test_ui_bt_queue_workers_controls_unique_and_wired():
    src = _v3_html()
    for eid in ("btQueueWorkers", "btQueueWorkersHint", "btQueueWorkersWrap"):
        assert src.count(f'id="{eid}"') == 1, eid
    assert '"/api/v1/backtests/queue/config"' in src
    fn = _extract_js_function(src, "btQueueWorkersOnChange")
    assert 'btQueueWorkersRequest("PUT"' in fn
    assert "max_workers: next" in fn
    assert "btQueueWorkersRequest" in _extract_js_function(src, "initBtQueueWorkers")


def test_ui_default_date_prefers_data_max_date_and_no_hardcoded_calendar_end():
    src = _v3_html()
    fn = _extract_js_function(src, "applyDefaultDateRange")
    assert "data_max_date" in fn
    # 优先 data_max_date，max_date 仅作回退（限定属性访问，避免子串误判）
    assert fn.index("__btCalMeta.data_max_date") < fn.index("__btCalMeta.max_date")

    load_fn = _extract_js_function(src, "loadBtCalendarAndDates")
    assert "data_max_date" in load_fn
    assert "20251231" not in load_fn, "兜底不得硬编码 20251231"
    assert "todayYmd" in load_fn


def test_ui_bq_hint_review_fallback_text_function_exists_and_used():
    src = _v3_html()
    assert src.count('id="bqHint"') == 1
    fn = _extract_js_function(src, "bqExportReviewSuffix")
    assert "review_asof_used" in fn
    assert "最后可用日" in fn
    # 同步/异步两条导出路径都使用该函数更新提示
    assert "bqExportReviewSuffix" in _extract_js_function(src, "bqHandleExportResponse")
    assert "bqExportReviewSuffix" in _extract_js_function(src, "bqPollExportJob")
