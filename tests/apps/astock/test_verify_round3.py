# -*- coding: utf-8 -*-
"""独立端到端验证（第 3 轮新功能）：
A. 用户规则走真实导出链路（真实 registry/user_registry，仅 mock 数据面）
   - 命中/0 命中/多规则；error 与 no_go 的占位空 sheet
B. 规则 API：import / batch-validate / categories / PATCH 边界（TestClient）
C. 前端静态 + Node 功能检查（规则编辑/导入/批量/分类）

不依赖 coder 新用例的内部桩；不改业务代码。
"""
from __future__ import annotations

import datetime as _dt
import io
import json
import re
import shutil
import subprocess
from pathlib import Path

import pytest

import tests.apps.astock.conftest  # noqa: F401

from wtpy.apps.astock.config import AStockConfig, get_default_config
from wtpy.apps.astock.data.tdx_reader import DayBar
from wtpy.apps.astock.service.rules import RuleService

BACKEND_ROOT = Path(__file__).resolve().parents[3]
V3_HTML = (
    BACKEND_ROOT / "wtpy" / "apps" / "astock" / "web" / "static" / "index_v3.html"
)
BAGUA_JSON = (
    BACKEND_ROOT / "wtpy" / "apps" / "astock" / "bagua" / "bagua_384.json"
)

_DS_META_MOCK = {
    "dataset_id": "mock",
    "dataset_source": "tdxquant",
    "dataset_adjustment": "front",
    "dataset_status": "ready",
    "covers_asof": True,
    "candidate_datasets": 1,
}

HIT_FORMULA = "XG:C>MA(C,5);"
ALWAYS_TRUE_FORMULA = "XG:MA(C,5)>0;"
ZERO_HIT_FORMULA = "XG:C>MA(C,5)*10000;"
MIN60_FORMULA = 'DIF60:="MACD.DIF#MIN60";\nXG:C>0 AND DIF60>0;'
MIN1_FORMULA = 'DIF1:="MACD.DIF#MIN1";\nXG:C>0 AND DIF1>0;'

ASOF = 20260828


# ---------------------------------------------------------------------------
# 公共
# ---------------------------------------------------------------------------


def _weekdays(start: str, end: str) -> list:
    out = []
    cur = _dt.date.fromisoformat(start)
    end_d = _dt.date.fromisoformat(end)
    while cur <= end_d:
        if cur.isoweekday() <= 5:
            out.append(int(cur.strftime("%Y%m%d")))
        cur += _dt.timedelta(days=1)
    return out


def _rising_bars() -> list:
    """稳定上涨收盘：末日 close 必大于 MA5（命中 HIT/ALWAYS_TRUE）。"""
    days = [d for d in _weekdays("2026-06-01", "2026-08-28") if d <= ASOF]
    closes = [10.0 + i * 0.05 for i in range(len(days))]
    return [DayBar(d, c, c + 0.01, c - 0.01, c, 1e6, 1e7) for d, c in zip(days, closes)]


def _cfg(tmp_path: Path) -> AStockConfig:
    cfg = get_default_config(
        storage_root=tmp_path / "st",
        indicator_dir=tmp_path / "ind",
        output_root=tmp_path / "out",
    )
    Path(cfg.storage_root).mkdir(parents=True, exist_ok=True)
    Path(cfg.indicator_dir).mkdir(parents=True, exist_ok=True)
    return cfg


def _mock_export_data(monkeypatch, cfg: AStockConfig, bars: list) -> None:
    """只 mock 数据面/卦象面；registry 与 user_registry 全走真实实现。"""
    from wtpy.apps.astock.service import bagua_query as bq

    monkeypatch.setattr(
        bq, "_load_dataset_bars", lambda *_a, **_k: (bars, dict(_DS_META_MOCK))
    )
    monkeypatch.setattr(
        bq,
        "BaguaPlaneSession",
        lambda *_a, **_k: (_ for _ in ()).throw(FileNotFoundError("no md")),
    )
    monkeypatch.setattr(
        bq,
        "_resolve_batch_codes",
        lambda cfg_, codes=None, *, all_stocks=False: ["SSE.STK.600000"],
    )
    monkeypatch.setattr(bq, "list_etf_std_codes", lambda cfg_: [])


def _mock_review_data(monkeypatch, cfg: AStockConfig, bars: list) -> None:
    """即时计算的数据面：正式 L1 表面 + 默认 bar loader。registry 不碰。"""
    from wtpy.apps.astock.service import indicator_review as ir

    monkeypatch.setattr(
        ir,
        "_resolve_formal_surface",
        lambda _cfg: ({"formal_l1_id": "mock_l1", "max_date": ASOF}, ""),
    )
    monkeypatch.setattr(
        ir,
        "_default_bar_loader_factory",
        lambda _cfg: (
            lambda code, asof: (
                [b for b in bars if int(b.date) <= int(asof)],
                {"dataset_id": "mock"},
            )
        ),
    )


def _export(cfg, selected, *, date="2026-08-28"):
    from wtpy.apps.astock.service import bagua_query as bq

    info: dict = {}
    path = bq.export_bagua_multi_period_xlsx(
        cfg,
        date=date,
        periods=["WEEK", "MONTH"],
        adjust="tushare_qfq",
        all_stocks=True,
        review_rules=selected,
        info_out=info,
    )
    return path, info


def _meta(path) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(path)
    return {r[0]: r[1] for r in wb["meta"].iter_rows(min_row=2, values_only=True)}


def _write_review(cfg: AStockConfig, asof: int, rules: list) -> None:
    d = Path(cfg.storage_root) / "indicator_review"
    d.mkdir(parents=True, exist_ok=True)
    payload = {
        "asof": int(asof),
        "generated_at": "2026-08-28 19:00:00",
        "status": "ok",
        "no_go_reason": "",
        "universe_size": 1,
        "scanned": 1,
        "error_count": 0,
        "rules": rules,
    }
    (d / f"review_{asof}.json").write_text(
        json.dumps(payload, ensure_ascii=False), encoding="utf-8"
    )


def _make_app(cfg):
    from wtpy.apps.astock.api import create_app

    return create_app(cfg)


def _make_client(app):
    from fastapi.testclient import TestClient

    return TestClient(app)


@pytest.fixture()
def rules_client(tmp_path: Path):
    pytest.importorskip("fastapi")
    cfg = _cfg(tmp_path)
    app = _make_app(cfg)
    try:
        yield _make_client(app), cfg
    finally:
        app.state.astock.jobs.shutdown(wait=False)


# ===========================================================================
# A. 导出：用户规则真实链路
# ===========================================================================


def test_round3_user_rule_hit_and_zero_hit_sheets_real_registry(
    tmp_path: Path, monkeypatch
):
    """RuleService 建规 → 真实导出（不 mock registry/bootstrap）：
    命中规则 sheet 有成员；0 命中规则 sheet 仅表头。"""
    if not BAGUA_JSON.exists():
        pytest.skip("bagua_384.json missing")

    cfg = _cfg(tmp_path)
    bars = _rising_bars()
    _mock_export_data(monkeypatch, cfg, bars)
    _mock_review_data(monkeypatch, cfg, bars)

    svc = RuleService(cfg)
    hit = svc.create_rule(name="独立趋势MA5", formula_text=HIT_FORMULA)
    zero = svc.create_rule(name="独立全假", formula_text=ZERO_HIT_FORMULA)
    assert hit["id"].startswith("user_") and zero["id"].startswith("user_")

    path, info = _export(cfg, [hit["id"], zero["id"]])

    import openpyxl

    wb = openpyxl.load_workbook(path)
    assert hit["name"] in wb.sheetnames, wb.sheetnames
    assert zero["name"] in wb.sheetnames, wb.sheetnames
    hit_rows = [r[0] for r in wb[hit["name"]].iter_rows(min_row=2, values_only=True)]
    assert hit_rows == ["600000"], "命中成员必须落表"
    assert wb[zero["name"]].max_row == 1, "0 命中仍是表头空 sheet"

    meta = _meta(path)
    sources = str(meta["indicator_review_rule_sources"])
    assert f"{hit['id']}=computed:{ASOF}" in sources
    assert f"{zero['id']}=computed:{ASOF}" in sources
    assert "placeholder" not in str(meta["indicator_review_note"])
    assert info["query_date"] == ASOF


def test_round3_multiple_user_rules_all_sheets_present(
    tmp_path: Path, monkeypatch
):
    """3 条 user 规则同时勾选（含重复传参）：全部出 sheet，不再一批全灭。"""
    if not BAGUA_JSON.exists():
        pytest.skip("bagua_384.json missing")

    cfg = _cfg(tmp_path)
    bars = _rising_bars()
    _mock_export_data(monkeypatch, cfg, bars)
    _mock_review_data(monkeypatch, cfg, bars)

    svc = RuleService(cfg)
    r1 = svc.create_rule(name="独立甲命中", formula_text=HIT_FORMULA)
    r2 = svc.create_rule(name="独立乙命中", formula_text=ALWAYS_TRUE_FORMULA)
    r3 = svc.create_rule(name="独立丙空", formula_text=ZERO_HIT_FORMULA)

    path, _info = _export(cfg, [r1["id"], r2["id"], r3["id"], r1["id"]])

    import openpyxl

    wb = openpyxl.load_workbook(path)
    signal = [n for n in wb.sheetnames if n not in ("meta", "stock-all")]
    assert signal == ["独立甲命中", "独立乙命中", "独立丙空"], wb.sheetnames
    assert [r[0] for r in wb["独立甲命中"].iter_rows(min_row=2, values_only=True)] == ["600000"]
    assert [r[0] for r in wb["独立乙命中"].iter_rows(min_row=2, values_only=True)] == ["600000"]
    assert wb["独立丙空"].max_row == 1

    meta = _meta(path)
    assert meta["indicator_review_sheets"] == "独立甲命中,独立乙命中,独立丙空"


def test_round3_compute_error_placeholder_coexists_with_precomputed(
    tmp_path: Path, monkeypatch
):
    """预计算命中照常出表 + 即时计算 error 的缺失规则补占位空表。"""
    if not BAGUA_JSON.exists():
        pytest.skip("bagua_384.json missing")

    cfg = _cfg(tmp_path)
    bars = _rising_bars()
    _mock_export_data(monkeypatch, cfg, bars)
    _write_review(
        cfg,
        ASOF,
        [
            {
                "rule_id": "txt_735金叉及趋势",
                "sheet": "735",
                "count": 1,
                "matched": [{"code": "SSE.STK.600000", "close": 5.9}],
            }
        ],
    )

    from wtpy.apps.astock.service import bagua_query as bq

    got_ids: list = []

    def _error_compute(cfg_, asof, rule_ids, *, codes=None, on_progress=None):
        got_ids.append(list(rule_ids))
        return {
            "asof": asof,
            "status": "error",
            "error_note": "即时计算失败（规则 user_broken_x）: KeyError",
            "rules": [],
        }

    monkeypatch.setattr(bq, "_compute_rules_for_export", _error_compute)

    path, _info = _export(cfg, ["txt_735金叉及趋势", "user_broken_x"])

    import openpyxl

    wb = openpyxl.load_workbook(path)
    assert [r[0] for r in wb["735"].iter_rows(min_row=2, values_only=True)] == ["600000"]
    assert "user_broken_x" in wb.sheetnames, wb.sheetnames
    assert wb["user_broken_x"].max_row == 1, "error 占位必须是仅表头空表"
    # 只把缺失的规则送给即时计算
    assert got_ids == [["user_broken_x"]]

    meta = _meta(path)
    note = str(meta["indicator_review_note"])
    assert "即时计算失败" in note
    assert "placeholder:user_broken_x(" in note
    assert "user_broken_x=" in str(meta["indicator_review_placeholders"])
    assert "user_broken_x=placeholder:" in str(meta["indicator_review_rule_sources"])
    assert meta["indicator_review_sheets"] == "735,user_broken_x"


def test_round3_compute_no_go_placeholder_for_all_missing(
    tmp_path: Path, monkeypatch
):
    """即时计算 no_go：全部未出 sheet 的已选规则补占位，reason=no_go 原因。"""
    if not BAGUA_JSON.exists():
        pytest.skip("bagua_384.json missing")

    cfg = _cfg(tmp_path)
    bars = _rising_bars()
    _mock_export_data(monkeypatch, cfg, bars)
    # 无复核文件 → 两条都进即时计算

    from wtpy.apps.astock.service import bagua_query as bq

    monkeypatch.setattr(
        bq,
        "_compute_rules_for_export",
        lambda cfg_, asof, rule_ids, *, codes=None, on_progress=None: {
            "asof": asof,
            "status": "no_go",
            "no_go_reason": "no_formal_l1_product",
            "rules": [],
        },
    )

    path, info = _export(cfg, ["user_alpha_ph", "user_beta_ph"])

    import openpyxl

    wb = openpyxl.load_workbook(path)
    for rid in ("user_alpha_ph", "user_beta_ph"):
        assert rid in wb.sheetnames, wb.sheetnames
        assert wb[rid].max_row == 1
    meta = _meta(path)
    note = str(meta["indicator_review_note"])
    assert "no_go:no_formal_l1_product" in note
    assert "placeholder:user_alpha_ph(no_formal_l1_product)" in note
    assert "placeholder:user_beta_ph(no_formal_l1_product)" in note
    ph = str(meta["indicator_review_placeholders"])
    assert "user_alpha_ph=no_formal_l1_product" in ph
    assert "user_beta_ph=no_formal_l1_product" in ph
    assert meta["indicator_review_sheets"] == "user_alpha_ph,user_beta_ph"
    assert info["query_date"] == ASOF


def test_round3_error_placeholder_uses_user_display_name(
    tmp_path: Path, monkeypatch
):
    """占位 sheet 名走真实 user_registry.json 的显示名（不是 user_ id）。"""
    if not BAGUA_JSON.exists():
        pytest.skip("bagua_384.json missing")

    cfg = _cfg(tmp_path)
    bars = _rising_bars()
    _mock_export_data(monkeypatch, cfg, bars)

    svc = RuleService(cfg)
    rule = svc.create_rule(name="占位显示名规则", formula_text=HIT_FORMULA)

    from wtpy.apps.astock.service import bagua_query as bq

    monkeypatch.setattr(
        bq,
        "_compute_rules_for_export",
        lambda cfg_, asof, rule_ids, *, codes=None, on_progress=None: {
            "asof": asof,
            "status": "error",
            "error_note": f"即时计算失败（规则 {rule_ids[0]}）: boom",
            "rules": [],
        },
    )

    path, _info = _export(cfg, [rule["id"]])

    import openpyxl

    wb = openpyxl.load_workbook(path)
    assert "占位显示名规则" in wb.sheetnames, wb.sheetnames
    assert wb["占位显示名规则"].max_row == 1
    meta = _meta(path)
    assert f"placeholder:{rule['id']}(" in str(meta["indicator_review_note"])
    assert rule["id"] not in wb.sheetnames


# ===========================================================================
# B. 规则 API 边界
# ===========================================================================

def test_round3_import_route_traversal_fields_and_guards(rules_client, tmp_path):
    client, cfg = rules_client
    ind = Path(cfg.indicator_dir)

    # 正常导入 + 列表/详情字段
    ok = client.post(
        "/api/v1/rules/import",
        json={"filename": "独立导入规则.txt", "content": HIT_FORMULA},
    )
    assert ok.status_code == 200, ok.text
    body = ok.json()
    assert body["id"].startswith("txt_") and body["name"] == "独立导入规则"
    row = next(r for r in client.get("/api/v1/rules").json() if r["id"] == body["id"])
    assert "description" in row and "category" in row
    assert row["description"] == "" and row["category"] == ""
    detail = client.get(f"/api/v1/rules/{body['id']}").json()
    assert "MA(C,5)" in detail["formula_text"]

    # 路径穿越：三种形态都必须取 basename 落在 indicator_dir
    cases = [
        ("..\\..\\escape.txt", "escape.txt"),
        ("/etc/passwd.txt", "passwd.txt"),
        ("C:\\Windows\\Temp\\hosts_verify.txt", "hosts_verify.txt"),
    ]
    for raw, base in cases:
        r = client.post(
            "/api/v1/rules/import", json={"filename": raw, "content": "XG:C>0;"}
        )
        assert r.status_code == 200, (raw, r.text)
        assert (ind / base).exists(), raw
    assert not (tmp_path / "escape.txt").exists()
    assert not (tmp_path.parent / "escape.txt").exists()

    # 后缀：.tn6 / .md 拒绝；大写 .TXT 接受
    assert client.post(
        "/api/v1/rules/import", json={"filename": "pkg3.tn6", "content": "x"}
    ).status_code == 400
    assert client.post(
        "/api/v1/rules/import", json={"filename": "note3.md", "content": "x"}
    ).status_code == 400
    assert client.post(
        "/api/v1/rules/import", json={"filename": "大写规则.TXT", "content": "XG:C>0;"}
    ).status_code == 200

    # 同名 400
    assert client.post(
        "/api/v1/rules/import",
        json={"filename": "独立导入规则.txt", "content": "XG:C>1;"},
    ).status_code == 400

    # 超 512KB / 空内容 / 空文件名：400 且不得落盘
    too_big = "x" * (512 * 1024 + 1)
    assert client.post(
        "/api/v1/rules/import", json={"filename": "big3.txt", "content": too_big}
    ).status_code == 400
    assert not (ind / "big3.txt").exists()
    assert client.post(
        "/api/v1/rules/import", json={"filename": "empty3.txt", "content": "   "}
    ).status_code == 400
    assert not (ind / "empty3.txt").exists()
    assert client.post(
        "/api/v1/rules/import", json={"filename": "", "content": "XG:C>0;"}
    ).status_code == 400


def test_round3_patch_route_formula_state_and_guards(rules_client):
    client, _cfg = rules_client

    created = client.post(
        "/api/v1/rules",
        json={
            "name": "独立回踩",
            "formula_text": "XG:C>0;",
            "description": "初版",
            "category": "独立趋势",
        },
    )
    assert created.status_code == 200, created.text
    rid = created.json()["id"]

    # 只改 name：公式状态原样
    ren = client.patch(f"/api/v1/rules/{rid}", json={"name": "独立回踩改"}).json()
    assert ren["name"] == "独立回踩改"
    assert ren["dependencies"] == [] and ren["min60_day_proxy"] is False

    # 与 create 同源：一条 create 直接给 MIN60，一条 PATCH 改成 MIN60，字段一致
    created_m60 = client.post(
        "/api/v1/rules", json={"name": "独立源对账m60", "formula_text": MIN60_FORMULA}
    ).json()
    patched_m60 = client.patch(
        f"/api/v1/rules/{rid}", json={"formula_text": MIN60_FORMULA}
    ).json()
    assert created_m60["dependencies"] == patched_m60["dependencies"] == ["MIN60"]
    assert created_m60["min60_day_proxy"] is True
    assert patched_m60["min60_day_proxy"] is True
    assert bool(created_m60["failure_reason"]) == bool(patched_m60["failure_reason"])
    assert created_m60["failure_reason"] and patched_m60["failure_reason"]

    # 去掉 MIN60：标记清除，name/description/category 保留
    back = client.patch(
        f"/api/v1/rules/{rid}",
        json={"formula_text": "XG:C>0;", "description": "改版", "category": "独立低吸"},
    ).json()
    assert back["dependencies"] == []
    assert back["min60_day_proxy"] is False
    assert not back["failure_reason"]
    assert back["name"] == "独立回踩改"
    assert back["description"] == "改版" and back["category"] == "独立低吸"

    # MIN1 拒绝且不污染已注册状态
    assert client.patch(
        f"/api/v1/rules/{rid}", json={"formula_text": MIN1_FORMULA}
    ).status_code == 400
    assert client.get(f"/api/v1/rules/{rid}").json()["dependencies"] == []

    # 非 user_ 规则 400；不存在的规则 404
    assert client.patch(
        "/api/v1/rules/bagua_ohlc", json={"name": "x"}
    ).status_code == 400
    assert client.patch(
        "/api/v1/rules/user_nope_123", json={"name": "x"}
    ).status_code == 404

    # 列表也带 description/category
    row = next(r for r in client.get("/api/v1/rules").json() if r["id"] == rid)
    assert row["description"] == "改版" and row["category"] == "独立低吸"


def test_round3_create_route_failure_leaves_no_artifacts(rules_client):
    """MIN1 / 语法错误创建失败：不留 user txt、不注册半截规则。"""
    client, cfg = rules_client
    user_dir = Path(cfg.storage_root) / "indicators" / "user"
    before_files = set(user_dir.glob("*.txt")) if user_dir.exists() else set()
    before_ids = {r["id"] for r in client.get("/api/v1/rules").json()}

    assert client.post(
        "/api/v1/rules", json={"name": "独立坏MIN1", "formula_text": MIN1_FORMULA}
    ).status_code == 400
    assert client.post(
        "/api/v1/rules", json={"name": "独立坏语法", "formula_text": "((("}
    ).status_code == 400
    assert client.post(
        "/api/v1/rules", json={"name": "独立无XG", "formula_text": "MA5:=MA(C,5);"}
    ).status_code == 400

    after_files = set(user_dir.glob("*.txt")) if user_dir.exists() else set()
    after_ids = {r["id"] for r in client.get("/api/v1/rules").json()}
    assert after_files == before_files, after_files - before_files
    assert after_ids == before_ids, after_ids - before_ids


def test_round3_batch_validate_route_dedup_and_partial_failure(rules_client):
    client, cfg = rules_client
    created = client.post(
        "/api/v1/rules", json={"name": "独立批量ok", "formula_text": "XG:C>0;"}
    ).json()
    ind = Path(cfg.indicator_dir)
    (ind / "broken_batch.txt").write_text("MA5:=MA(C,5);\n", encoding="utf-8")
    (ind / "pkg_batch.tn6").write_bytes(b"pkg")

    r = client.post(
        "/api/v1/rules/batch-validate",
        json={
            "ids": [
                created["id"],
                created["id"],
                "",
                "   ",
                "txt_broken_batch",
                "tn6_pkg_batch",
                "nope_x",
            ]
        },
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["summary"] == {"total": 4, "ok": 1, "failed": 3}
    by_id = {x["id"]: x for x in body["results"]}
    assert by_id[created["id"]]["ok"] is True
    assert by_id["txt_broken_batch"]["ok"] is False and by_id["txt_broken_batch"]["error"]
    assert by_id["tn6_pkg_batch"]["ok"] is False
    assert by_id["nope_x"]["ok"] is False

    # ids 缺省=全量，且 ok+failed 恒等于 total
    all_body = client.post("/api/v1/rules/batch-validate", json={}).json()
    s = all_body["summary"]
    assert s["total"] >= 4 and s["ok"] + s["failed"] == s["total"]


def test_round3_categories_route_roundtrip_restart_and_validation(
    rules_client, tmp_path
):
    client, cfg = rules_client
    assert client.get("/api/v1/rules/categories").json() == {"categories": []}

    assert client.post(
        "/api/v1/rules/categories", json={"name": "独立分类A"}
    ).json() == {"categories": ["独立分类A"]}
    # 前后空白：幂等
    assert client.post(
        "/api/v1/rules/categories", json={"name": "  独立分类A  "}
    ).json() == {"categories": ["独立分类A"]}
    assert client.post(
        "/api/v1/rules/categories", json={"name": "独立分类B"}
    ).json() == {"categories": ["独立分类A", "独立分类B"]}
    assert client.get("/api/v1/rules/categories").json() == {
        "categories": ["独立分类A", "独立分类B"]
    }
    assert (Path(cfg.storage_root) / "indicators" / "rule_categories.json").exists()

    # 重启新 app（同 cfg）后分类仍在
    app2 = _make_app(cfg)
    try:
        assert _make_client(app2).get("/api/v1/rules/categories").json() == {
            "categories": ["独立分类A", "独立分类B"]
        }
    finally:
        app2.state.astock.jobs.shutdown(wait=False)

    # 非法名：空 / 全空白 / 超 20 字 → 400；恰好 20 字可接受
    assert client.post(
        "/api/v1/rules/categories", json={"name": "   "}
    ).status_code == 400
    assert client.post(
        "/api/v1/rules/categories", json={"name": "x" * 21}
    ).status_code == 400
    assert client.post(
        "/api/v1/rules/categories", json={"name": "y" * 20}
    ).status_code == 200


# ===========================================================================
# C. 前端静态 / Node
# ===========================================================================


def _extract_js_function(src: str, name: str) -> str:
    m = re.search(r"(?:async\s+)?function\s+" + name + r"\s*\(", src)
    assert m, f"function not found: {name}"
    brace = src.index("{", m.start())
    depth = 0
    for j in range(brace, len(src)):
        if src[j] == "{":
            depth += 1
        elif src[j] == "}":
            depth -= 1
            if depth == 0:
                return src[m.start(): j + 1]
    raise AssertionError(f"unbalanced {name}")


def test_round3_ui_rule_editor_import_batch_category_wiring():
    assert V3_HTML.is_file()
    src = V3_HTML.read_text(encoding="utf-8")

    # 元素 id 唯一
    for eid in ("ruleImportFile", "ruleBatchResult", "newRuleCategory"):
        assert src.count(f'id="{eid}"') == 1, eid

    # 关键函数各定义一次
    for fn in (
        "openRuleEditor",
        "resetRuleForm",
        "onRuleImportFileChange",
        "runBatchRuleTest",
        "renderBatchRuleResult",
        "createRuleCategory",
        "ruleCategory",
    ):
        assert len(re.findall(r"(?:async\s+)?function\s+" + fn + r"\s*\(", src)) == 1, fn

    # 编辑：GET 回填 + PATCH 保存；新建/取消/保存后表单清空
    editor = _extract_js_function(src, "openRuleEditor")
    assert 'api("/api/v1/rules/"' in editor
    for eid in ("newRuleName", "newRuleFormula", "newRuleDesc", "newRuleCategory"):
        assert eid in editor, eid
    reset = _extract_js_function(src, "resetRuleForm")
    assert "__ruleEditId = null" in reset
    for eid in ("newRuleName", "newRuleFormula", "newRuleDesc", "newRuleCategory"):
        assert eid in reset, eid
    save = _extract_js_function(src, "saveNewRule")
    assert 'method: "PATCH"' in save and '"/api/v1/rules"' in save
    assert "resetRuleForm()" in save
    assert re.search(
        r'btnCancelCreate"\)\.onclick = \(\) => \{\s*resetRuleForm\(\)', src
    )
    assert re.search(
        r'btnNewRule"\)\.onclick = \(\) => \{\s*resetRuleForm\(\)', src
    )

    # 导入：文件选择 → JSON POST；批量：结果面板
    imp = _extract_js_function(src, "onRuleImportFileChange")
    assert "ruleImportFile" in imp
    assert 'api("/api/v1/rules/import"' in imp
    batch = _extract_js_function(src, "runBatchRuleTest")
    assert 'api("/api/v1/rules/batch-validate"' in batch
    panel = _extract_js_function(src, "renderBatchRuleResult")
    assert "ruleBatchResult" in panel and "summary" in panel

    # 分类：prompt → POST → AppState.ruleCategories / datalist
    cat = _extract_js_function(src, "createRuleCategory")
    assert "prompt(" in cat and 'api("/api/v1/rules/categories"' in cat
    assert "AppState.ruleCategories" in cat
    assert "renderRuleCenter()" in cat
    # 分类变更 → renderRuleCenter → refreshRuleCategoryOptions（表单 datalist 更新）
    assert "refreshRuleCategoryOptions()" in _extract_js_function(src, "renderRuleCenter")

    # ruleCategory 优先 rule.category（Node 真跑）
    node = shutil.which("node")
    assert node, "node required"
    fn = _extract_js_function(src, "ruleCategory")
    script = (
        fn
        + "\n"
        + """
function assert(cond, msg) { if (!cond) throw new Error(msg); }
assert(ruleCategory({category: " 我的分类 ", source: "user", name: "x"}) === "我的分类", "category priority");
assert(ruleCategory({source: "user", name: "x"}) === "自定义", "user fallback");
assert(ruleCategory({name: "735金叉"}) === "动量", "builtin momentum");
assert(ruleCategory({name: "卦象观察"}) === "卦象", "builtin gua");
assert(ruleCategory({name: "先跌后涨"}) === "反转", "builtin reversal");
assert(ruleCategory({}) === "内置", "default builtin");
console.log("PASS ruleCategory");
"""
    )
    proc = subprocess.run(
        [node, "-e", script],
        capture_output=True,
        text=True,
        cwd=str(BACKEND_ROOT),
    )
    assert proc.returncode == 0, proc.stdout + "\n" + proc.stderr
    assert "PASS ruleCategory" in proc.stdout
