# 一次性验证脚本：卦象导出是否真的按传入历史日期取周/月K线
# 用法：python -m tools.verify_export_history_date [--rizhu-path PATH]
import argparse
import sys
from pathlib import Path

from wtpy.apps.astock.config import get_default_config
from wtpy.apps.astock.service.bagua_query import export_bagua_multi_period_xlsx


def _parse_args(argv=None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="验证卦象导出是否按传入历史日期取周/月K线",
    )
    parser.add_argument(
        "--rizhu-path",
        default=None,
        help="日柱表 xlsx 路径（可选；不传时按上市日期推算日柱）",
    )
    return parser.parse_args(argv)


def main(argv=None) -> int:
    args = _parse_args(argv)
    cfg = get_default_config()
    rizhu = Path(args.rizhu_path) if args.rizhu_path else None
    if rizhu is not None and not rizhu.exists():
        print(f"错误：--rizhu-path 指定的文件不存在：{rizhu}", file=sys.stderr)
        print(
            "提示：不传该参数时脚本跳过日柱表，按上市日期推算日柱。",
            file=sys.stderr,
        )
        return 2
    if rizhu is None:
        print("提示：未指定 --rizhu-path，日柱将按上市日期推算（不读取日柱表）。")

    # 历史日期：2024-01-28 是周日（ISO 2024-W04 的周日）
    hist_date = "2024-01-28"
    out = export_bagua_multi_period_xlsx(
        cfg,
        date=hist_date,
        periods=["WEEK", "MONTH"],
        adjust="tushare_qfq",
        codes=["000001", "600000", "000333"],
        all_stocks=False,
        rizhu_path=rizhu,
    )
    print("exported:", out)

    import openpyxl
    wb = openpyxl.load_workbook(out, read_only=True)
    print("sheets:", wb.sheetnames)
    ws = wb["stock-all"] if "stock-all" in wb.sheetnames else wb[wb.sheetnames[-1]]
    rows = list(ws.iter_rows(values_only=True))
    headers = rows[0]
    # 定位关键列
    idx = {h: i for i, h in enumerate(headers)}
    week_end_i = idx.get("week_end")
    week_combo_i = next((i for h, i in idx.items() if h.startswith("周卦周线-组合")), None)
    month_combo_i = next((i for h, i in idx.items() if h.startswith("月卦月线-组合")), None)
    status_i = idx.get("数据状态")
    print("headers:", headers)
    print()
    for r in rows[1:]:
        print(
            "code=", r[idx["code"]],
            "| week_end=", r[week_end_i],
            "| 周卦=", r[week_combo_i],
            "| 月卦=", r[month_combo_i],
            "| status=", r[status_i],
        )
    print()
    print(">>> 周卦表头标签:", headers[week_combo_i] if week_combo_i is not None else "?")
    print(">>> 月卦表头标签:", headers[month_combo_i] if month_combo_i is not None else "?")
    print(">>> 传入历史 date:", hist_date, "(2024-01-28, 周日, ISO 2024-W04)")
    print(">>> 期望: week_end 落在 2024-01-22 ~ 2024-01-28; 月卦标签为 2023-12")
    return 0


if __name__ == "__main__":
    sys.exit(main())
